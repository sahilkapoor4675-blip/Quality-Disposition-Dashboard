#!/usr/bin/env python3
"""Excel / PDF / PPTX export checks.

Merged from export_acceptance.py + export_stress.py.
  python tests/test_exports.py            # acceptance only (fast, ~6 s)
  python tests/test_exports.py --stress   # acceptance + high-cardinality stress (~35 s)
"""
import io, os, shutil, subprocess, sys, tempfile, time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent   # repo root (this file lives in tests/)
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
os.environ['APP_VERSION'] = 'test'

from openpyxl import load_workbook
from pptx import Presentation
from reports import _excel_report, _pdf_report, _pptx_report


def export_acceptance():
    N=25

    def rows(prefix):
        return [{'name':f'{prefix}_{i}','output_qty':i+1,'coils':i+2,'defect_coils':1,'defect_pct':0.01,'reject_qty':0.1,'reject_pct_qty':0.001} for i in range(N)]

    wc=rows('WC'); gr=rows('GR')
    trend=[{'name':f'P{i}','coils':10,'output_qty':10,'defect_coils':1,'defect_pct':.1,'reject_qty':.1,'reject_pct_qty':.01,'first_pass_yield_pct':.9} for i in range(N)]
    reg=[{'rank':i+1,'defect':f'D{i}','records':1,'qty':float(i+1),'pct_records':1/N,'cum_pct':min(1,(i+1)/N)} for i in range(N)]
    causes={k:[f'Cause {i}' for i in range(8)] for k in ('man','machine','material','method','measurement','environment')}
    rca={k:[{'root_cause':f'RC{i}','action':'Action','preventive_action':'Prevent','role':'Role','responsibility':'Owner','why_chain':['Why 1','Why 2']} for i in range(3)] for k in causes}
    payload={
     'filters': {'month':'All','week':'All','quarter':'All','financial_year':'All','work_center':'All','grade':'All','main_defect':'All','defect_intensity':'All','quality_decision':'All'},
     'kpis': {'kpis':[{'label':f'K{i}','value':i,'fmt':''} for i in range(12)],'decision_table':[{'decision':f'DEC_{i}','qty':i+1} for i in range(6)],'intensity_table':[{'intensity':f'I{i}','qty':i+1,'coils':10} for i in range(8)]},
     'defects': {'register':reg,'pareto':reg[:10],'register_total':{'records':N,'qty':sum(range(1,N+1)),'pct_records':1.0}},
     'wcg': {'by_work_center':wc,'by_grade':gr,'total_work_center':{'coils':999,'output_qty':999,'defect_coils':1,'defect_pct':.01,'reject_qty':.1,'reject_pct_qty':.001},'total_grade':{'coils':999,'output_qty':999,'defect_coils':1,'defect_pct':.01,'reject_qty':.1,'reject_pct_qty':.001}},
     'monthly': {'rows':trend,'total':{}}, 'period': {'rows':trend,'total':{}}, 'quarterly': {'rows':trend[:10],'total':{}}, 'yearly': {'rows':trend[:5],'total':{}},
     'intel': {'early_warnings':[],'recurring_patterns':[],'kpi_ranking':[],'risk_matrix':{'work_centers':wc,'grades':gr},'health_score':{'reasons':[]}},
     'root_cause': {'defect':'D0','rows':[{'grade':'A','work_center':'WC','heat_no':'H','batch_no':'B','output_weight':1.0} for _ in range(N)]},
     'target_history': {'rows':trend}, 'fishbone': {'matched':True,'defect':'D0','causes':causes,'rca':rca}, 'fishbone_style':{}
    }

    with tempfile.TemporaryDirectory() as td:
        out=Path(td)
        xlsx=out/'report.xlsx'; pdf=out/'report.pdf'; pptx=out/'report.pptx'
        xlsx.write_bytes(_excel_report(payload)); pdf.write_bytes(_pdf_report(payload)); pptx.write_bytes(_pptx_report(payload))
        wb=load_workbook(xlsx,read_only=True,data_only=False)
        expected={'Defect Analysis':N+4,'Work Center':N+4,'Grade Analysis':N+4,'Monthly Trend':N+3,'Weekly Trend':N+3,'Quarterly Trend':12+1,'Financial Year':7+1,'6M Fishbone Analysis':1+1+1+N*0+48,'Target vs Actual History':N+3}
        for sheet, min_rows in expected.items():
            assert wb[sheet].max_row >= min_rows, f'{sheet} appears truncated: {wb[sheet].max_row}'
        prs=Presentation(pptx)
        required=('Quality Decision Distribution','Defect Analysis','Work Center Performance','Grade Performance','Monthly Trend','Weekly Trend','Quarterly Trend','Financial Year Trend')
        hits={k:0 for k in required}
        for slide in prs.slides:
            text=' | '.join(sh.text for sh in slide.shapes if hasattr(sh,'text'))
            pics=sum(1 for sh in slide.shapes if getattr(sh,'shape_type',None)==13)
            tables=sum(1 for sh in slide.shapes if getattr(sh,'has_table',False))
            for k in required:
                if k in text:
                    hits[k]+=1
                    assert pics >= 1, f'{k} slide missing chart image'
                    assert tables >= 1, f'{k} slide missing paired table'
        for k,v in hits.items():
            assert v >= 1, f'{k} missing from PPT'
    print('EXPORT ACCEPTANCE PASS')


def export_stress():
    def mk_rows(n, prefix):
        return [{'name':f'{prefix}_{i}','output_qty':float(i+1),'coils':1,'defect_coils':0,'defect_pct':0.0,'reject_qty':0.0,'reject_pct_qty':0.0} for i in range(n)]

    N=2000
    wc=mk_rows(N,'WC'); gr=mk_rows(N,'GR')
    trend=[{'name':f'P{i}','coils':100,'output_qty':100.0,'defect_coils':5,'defect_pct':0.05,'reject_qty':3.0,'reject_pct_qty':0.03,'first_pass_yield_pct':0.92} for i in range(120)]
    reg=[{'rank':i+1,'defect':f'D{i}','records':1,'qty':1.0,'pct_records':0.001,'cum_pct':min(1,(i+1)/10)} for i in range(N)]
    causes={k:[f'Cause {i}' for i in range(200)] for k in ('man','machine','material','method','measurement','environment')}
    rca={k:[{'root_cause':f'RC{i}','action':'Action','preventive_action':'Prevent','role':'Role','responsibility':'Owner'} for i in range(100)] for k in causes}

    payload={
     'filters': {'month':'All','week':'All','quarter':'All','financial_year':'All','work_center':'All','grade':'All','main_defect':'All','defect_intensity':'All','quality_decision':'All'},
     'kpis': {'kpis':[{'label':f'K{i}','value':i,'fmt':''} for i in range(12)],'decision_table':[{'decision':'PRIME','qty':100.0}],'intensity_table':[{'intensity':'LIGHT','qty':100.0,'coils':10}]},
     'defects': {'register':reg,'pareto':reg[:10],'register_total':{'records':N,'qty':float(N),'pct_records':1.0}},
     'wcg': {'by_work_center':wc,'by_grade':gr,'total_work_center':{},'total_grade':{}},
     'monthly': {'rows':trend[:100],'total':{}}, 'period': {'rows':trend[:100],'total':{}}, 'quarterly': {'rows':trend[:40],'total':{}}, 'yearly': {'rows':trend[:20],'total':{}},
     'intel': {'early_warnings':[{'title':'Warning','detail':'High cardinality stress','action':'Review'}]*60,'recurring_patterns':[{'defect':'D','grade':'G','work_center':'W','period_count':5,'qty':1.0}]*100,'kpi_ranking':[],'risk_matrix':{'work_centers':wc,'grades':gr},'health_score':{'reasons':[]}},
     'root_cause': {'defect':'D0','rows':[{'grade':'A','work_center':'WC','heat_no':'H','batch_no':'B','output_weight':1.0} for _ in range(1000)]},
     'target_history': {'rows':trend[:100]}, 'fishbone': {'matched':True,'defect':'D0','causes':causes,'rca':rca}, 'fishbone_style':{}
    }

    import io, shutil, subprocess, tempfile
    outputs={}
    for name, fn in [('excel', _excel_report), ('pdf', _pdf_report), ('pptx', _pptx_report)]:
        t=time.time(); data=fn(payload); assert data and len(data)>1000
        outputs[name]=data
        print(f'{name}_export=PASS bytes={len(data)} sec={time.time()-t:.2f}')

    # Completeness: high-cardinality tables must be printed IN FULL (no top-N cap), only chart visuals are bounded.
    import openpyxl
    wb=openpyxl.load_workbook(io.BytesIO(outputs['excel']), read_only=True)
    for sheet in ('Defect Analysis','Work Center','Grade Analysis'):
        n=sum(1 for _ in wb[sheet].iter_rows(values_only=True))
        assert n>=N, f'Excel sheet {sheet} has only {n} rows (< {N})'
    from pptx import Presentation
    prs=Presentation(io.BytesIO(outputs['pptx']))
    totals={}
    for slide in prs.slides:
        title=next((sh.text_frame.text.strip().split('\n')[0] for sh in slide.shapes if getattr(sh,'has_text_frame',False) and sh.text_frame.text.strip()), None)
        for sh in slide.shapes:
            if getattr(sh,'has_table',False) and sh.has_table:
                totals[title]=totals.get(title,0)+len(sh.table.rows)-1
    for title in ('Defect Analysis','Work Center Performance','Grade Performance'):
        assert totals.get(title,0)>=N, f'PPTX "{title}" lists only {totals.get(title,0)} rows (< {N})'
    if shutil.which('pdftotext'):
        fp=tempfile.mktemp(suffix='.pdf'); open(fp,'wb').write(outputs['pdf'])
        text=subprocess.run(['pdftotext','-layout',fp,'-'],capture_output=True,text=True).stdout
        for tag in (f'D{N-1}', f'WC_{N-1}', f'GR_{N-1}'):
            assert tag in text, f'PDF is missing the last row {tag} (table truncated)'
        assert 'Showing top' not in text, 'PDF still contains a top-N truncation note'
    print('EXPORT STRESS PASS')


def table_export_acceptance():
    """14th pass: the generic styled table workbook (drill-down + Chemistry exports) keeps every row and column, colours decisions and neutralises formulas."""
    import io, openpyxl, reports
    cols=[{'label':'Date','icon':'📅','kind':'date'},{'label':'Heat','icon':'🔥','kind':'heat'},{'label':'MT','kind':'num3'},{'label':'Decision','kind':'decision'},{'label':'Status','kind':'status'},{'label':'Cpk','kind':'idx','thr':{'target':1.33,'warning':1.0,'direction':'higher'}}]
    rows=[['12-04-2026',f'NBS{i}',1.5,'REJECT' if i%2 else 'PRIME','OUT OF SPEC · ABOVE USL' if i%2 else 'OK',0.8 if i%2 else 1.6] for i in range(1200)]
    rows[0][1]='=HYPERLINK("http://x")'
    data=reports._table_xlsx({'title':'T','sections':[{'title':'Records','columns':cols,'rows':rows,'total':['Total']}]})
    ws=openpyxl.load_workbook(io.BytesIO(data)).active
    body=[r for r in ws.iter_rows(values_only=True) if r[1] and str(r[1]).lstrip("'").startswith(('NBS','=HYPER'))]
    assert len(body)==1200, f'every row must be exported (got {len(body)})'
    assert str(body[0][1]).startswith("'="), 'formula injection must be neutralised'
    assert len(ws._images)==1, 'JSL logo is on the sheet'
    fills={c.value:c.fill.fgColor.rgb for r in ws.iter_rows() for c in r if c.value in ('REJECT','PRIME')}
    assert fills['REJECT']!=fills['PRIME'], 'decisions are coloured differently'
    print('table export OK (1200 rows, logo, colours, formula-safe)')


if __name__ == "__main__":
    export_acceptance()
    table_export_acceptance()
    if "--stress" in sys.argv:
        export_stress()
