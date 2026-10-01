#!/usr/bin/env python3
"""Real-browser regression checks (Chromium via Playwright).

The other gate checks (regression.py, static_checks.py, test_smoke.py) run the
server and read source files, but none of them actually load a page in a
browser -- so a bug that only shows up in the DOM/JS runtime can pass every
other check and still ship. Three bugs found in manual testing were exactly
that kind: the admin sidebar highlighting the wrong section on click, filters
letting you pick a combination with zero overlapping records, and the CSP
tightening in server.py silently breaking an inline onclick="..." handler.
This file drives a real headless Chromium against a real server process and
checks the actual rendered/runtime behaviour, so those regressions fail the
gate instead of reaching production.

Requires the optional `playwright` package with the Chromium browser
installed (`pip install playwright && playwright install chromium`). This is
a dev/CI dependency, not a runtime one -- it is never imported by server.py
-- so run_gate.py adds this check only when both are available, and skips it
(with a visible note, not a silent pass) otherwise.

    python tests/test_browser.py
"""
import os, shutil, subprocess, sys, tempfile, time
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
PORT = int(os.environ.get("BROWSER_TEST_PORT", "8799"))
BASE = f"http://127.0.0.1:{PORT}"
ADMIN_USER = "tester"   # NOT "admin": the users table intentionally hides the
                        # reset-password/toggle actions for that literal username,
                        # so a different seed name is needed to test those buttons.
ADMIN_PASS = "Strong-Admin-1234!"


def _start_server(db_path):
    env = os.environ.copy()
    env["DB_PATH"] = str(db_path)          # isolated, disposable copy -- never the real DB
    env["BACKUP_DIR"] = str(db_path.parent / "backups")
    env["ADMIN_USERNAME"] = ADMIN_USER
    env["ADMIN_PASSWORD"] = ADMIN_PASS
    env["APP_VERSION"] = "test"
    env.pop("RENDER", None)
    proc = subprocess.Popen([sys.executable, str(ROOT / "server.py"), str(PORT)], cwd=ROOT,
                             stdout=subprocess.PIPE, stderr=subprocess.STDOUT, env=env, text=True)
    import urllib.request
    for _ in range(60):
        try:
            urllib.request.urlopen(BASE + "/api/health", timeout=2)
            return proc
        except Exception:
            if proc.poll() is not None:
                out = proc.stdout.read(3000) if proc.stdout else ""
                raise AssertionError("server exited before becoming ready:\n" + out)
            time.sleep(0.2)
    proc.terminate()
    raise AssertionError("server did not become ready in time")


def _seed_chemistry():
    """Import a Standard.xlsx (Standard + AIM sheets) and a cast-chemistry workbook through the
    real admin API, so the Chemistry SPC tab has data to render. The bundled quality.db carries
    only disposition rows; without this the tab (correctly) shows its "No chemistry data" state
    and none of the chemistry UI assertions below could ever pass."""
    import io, json, re, urllib.request
    import openpyxl
    sys.path.insert(0, str(ROOT))

    def xlsx_bytes(wb):
        bio = io.BytesIO(); wb.save(bio); return bio.getvalue()

    hdr = ["Alloy", "Grade Descriptions", "Cu% (LSL)", "Cu% (USL)", "Zn% (LSL)", "Zn% (USL)", "Ni% (LSL)", "Ni% (USL)", "Pb% (LSL)", "Pb% (USL)"]
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Standard"; ws.append(hdr)
    ws.append(["NBS", "Test Brass", 74, 76, 19, 21, None, None, 0, .04])
    ws.append(["CNI", "Cu-Ni 80-20", 78, 82, None, None, 19, 21, 0, .03])
    w2 = wb.create_sheet("AIM"); w2.append(hdr)
    w2.append(["NBS", "Test Brass", 74.2, 75.8, 19.2, 20.8, None, None, 0, .01])
    w2.append(["CNI", "Cu-Ni 80-20", 78.5, 81.5, None, None, 19.5, 20.5, 0, .02])
    std = xlsx_bytes(wb)

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Test Brass"
    ws.append(["Date", "Coil No.", "Alloy", "Denomination", "Cu%", "Ni%", "Zn%", "Pb%", "Total%", "Name"])
    for i in range(30):
        cu = 75.0 + (i % 5 - 2) * 0.05; zn = 20 - (i % 5 - 2) * 0.05
        ws.append([f"{1 + i % 27:02d}.05.2026", "NBS6348" if i == 0 else f"NBS9{i:03d}", "NBS", "5RS",
                   cu, 5.0, zn, .003, cu + 5.0 + zn + .003, "SUNIL"])
    # Extra Test Brass heats in March 2026 (FY 2025-26, Q4) and June 2026 (FY 2026-27, Q1) next to the 30 May heats,
    # plus a second grade: enough variety for Month / Week / Quarter / Fin. Year to cascade for real.
    import datetime as _dt
    for i in range(8):
        d = _dt.date(2026, 3, 2) + _dt.timedelta(days=3 * i); cu = 75.0 + (i % 3 - 1) * 0.1; zn = 20 - (i % 3 - 1) * 0.1
        ws.append([d.strftime("%d.%m.%Y"), f"NBS7{100 + i}", "NBS", "5RS", cu, 5.0, zn, .003, cu + 5.0 + zn + .003, "SUNIL"])
    for i in range(6):
        d = _dt.date(2026, 6, 3) + _dt.timedelta(days=4 * i); cu = 75.0 + (i % 3 - 1) * 0.1; zn = 20 - (i % 3 - 1) * 0.1
        ws.append([d.strftime("%d.%m.%Y"), f"NBS7{200 + i}", "NBS", "5RS", cu, 5.0, zn, .003, cu + 5.0 + zn + .003, "SUNIL"])
    # Each grade has its own sheet, like the real files (a sheet named exactly like a grade is matched to that grade first).
    ws2 = wb.create_sheet("Cu-Ni 80-20")
    ws2.append(["Date", "Coil No.", "Alloy", "Denomination", "Cu%", "Ni%", "Zn%", "Pb%", "Total%", "Name"])
    for i in range(10):
        d = _dt.date(2026, 5, 4) + _dt.timedelta(days=3 * i); cu = 80.0 + (i % 4 - 1.5) * 0.2; ni = 20.0 - (i % 4 - 1.5) * 0.15
        ws2.append([d.strftime("%d.%m.%Y"), f"CNI8{100 + i}", "CNI", "X", cu, ni, 0, .002, cu + ni + .002, "SUNIL"])
    chem = xlsx_bytes(wb)

    req = urllib.request.Request(BASE + "/api/login", method="POST",
                                 data=json.dumps({"username": ADMIN_USER, "password": ADMIN_PASS}).encode(),
                                 headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=10) as r:
        cookie = "; ".join(h.split(";")[0] for h in r.headers.get_all("Set-Cookie"))
    csrf = re.search(r"qdash_csrf=([^;]+)", cookie).group(1)

    def call(path, body=None, file=None):
        h = {"Cookie": cookie, "X-CSRF-Token": csrf}
        if file:
            b = "----seed" + "x" * 12
            data = (f'--{b}\r\nContent-Disposition: form-data; name="file"; filename="{file[0]}"\r\n'
                    f'Content-Type: application/octet-stream\r\n\r\n').encode() + file[1] + f"\r\n--{b}--\r\n".encode()
            h["Content-Type"] = f"multipart/form-data; boundary={b}"
        else:
            data = json.dumps(body).encode(); h["Content-Type"] = "application/json"
        with urllib.request.urlopen(urllib.request.Request(BASE + path, data=data, headers=h, method="POST"), timeout=30) as r:
            return json.loads(r.read().decode())

    pv = call("/api/admin/chem_specs_preview", file=("Standard.xlsx", std))
    call("/api/admin/chem_specs_confirm", {"preview_id": pv["preview_id"]})
    pv = call("/api/admin/chem_import_preview", file=("chem.xlsx", chem))
    done = call("/api/admin/chem_import_confirm", {"preview_id": pv["preview_id"]})
    assert done.get("inserted") == 54, f"chemistry seed failed: {done}"


def run():
    from playwright.sync_api import sync_playwright

    tmpdir = Path(tempfile.mkdtemp(prefix="qdash_browser_"))
    db_path = tmpdir / "quality.db"
    shutil.copy2(ROOT / "quality.db", db_path)
    proc = _start_server(db_path)
    violations = []   # CSP violations, collected across the WHOLE run (every page)
    console_errors = []
    page_errors = []
    try:
        # Seeding is inside the try on purpose: if it fails, the finally below still stops the server
        # (a leaked server on this port made the NEXT run talk to a stale, already-seeded instance).
        _seed_chemistry()
        with sync_playwright() as p:
            browser = p.chromium.launch(args=["--no-sandbox"])
            context = browser.new_context(viewport={"width": 1440, "height": 900})
            # The real Google Fonts CDN is deliberately unreachable in CI; abort it so
            # the test can't flake on outbound network access it doesn't actually need.
            context.route("**/fonts.g*/**", lambda route: route.abort())

            def watch(page, label):
                page.add_init_script(
                    "document.addEventListener('securitypolicyviolation', e => "
                    "(window.__csp = window.__csp || []).push(e.violatedDirective + ': ' + e.blockedURI));"
                )
                page.on("console", lambda m: console_errors.append(f"{label}: {m.text[:200]}")
                        if m.type == "error" and "ERR_FAILED" not in m.text else None)
                page.on("pageerror", lambda e: page_errors.append(f"{label}: {str(e)[:300]}"))

            def collect_csp(page, label):
                for v in page.evaluate("window.__csp || []"):
                    violations.append(f"{label}: {v}")

            # ---- Dashboard: load, tabs, filter cascade, empty-selection state ----
            page = context.new_page()
            watch(page, "dashboard")
            page.goto(BASE + "/", wait_until="load")
            page.wait_for_timeout(1000)
            page.click("#introDashboardBtn")
            page.wait_for_timeout(2000)
            assert page.locator(".kpi-card").count() >= 8, "dashboard KPI cards did not render"

            for i in range(1, 6):
                page.click(f"#tabs button:nth-child({i})")
                page.wait_for_timeout(600)
            page.click("#tabs button:nth-child(1)")
            page.wait_for_timeout(400)

            # Cascading filters (#4): narrowing one dropdown must narrow the others.
            def options_for(key):
                return page.eval_on_selector_all(
                    f'.filter-field[data-filter-key="{key}"] .filter-options > *',
                    "els => els.map(e => e.textContent)")
            months_before = options_for("month")
            page.click('.filter-field[data-filter-key="quarter"] .filter-trigger')
            page.wait_for_timeout(200)
            page.click('.filter-field[data-filter-key="quarter"] .filter-option[data-value="Q2"]')
            page.wait_for_timeout(1200)
            months_after = options_for("month")
            assert len(months_after) < len(months_before), (
                f"filters did not cascade: month options unchanged after Quarter=Q2 "
                f"({months_before} -> {months_after})")

            # Empty-selection state: an impossible combo must say so, not show fake 0% KPIs.
            page.click('.filter-field[data-filter-key="month"] .filter-trigger')
            page.wait_for_timeout(200)
            # Pick a month NOT in the (already narrowed) list, forcing a zero-overlap combo,
            # by re-selecting a month from the ORIGINAL full list that quarter=Q2 excludes.
            excluded_month = next(m for m in months_before if m not in months_after and m != "All")
            page.evaluate("()=>document.querySelectorAll('.filter-field[data-filter-key=\"month\"] .filter-trigger')[0].click()")
            page.wait_for_timeout(150)
            # The option for excluded_month no longer exists in the (correctly) narrowed
            # dropdown -- which is the #4 fix working. To still exercise the #2 empty-state
            # banner, force the combination directly via the URL (a bookmarked/shared link
            # can still encode a stale combo even though the UI itself now prevents picking one).
            page.keyboard.press("Escape")
            page.goto(f"{BASE}/?month={excluded_month.replace(' ', '%20')}&quarter=Q2")
            # A full reload replays the intro splash (it is never skipped by design), so it
            # must be acknowledged again before anything behind it can be clicked.
            page.wait_for_timeout(1000)
            page.click("#introDashboardBtn")
            page.wait_for_timeout(1800)
            banner = page.inner_text("#periodBanner")
            assert "No records match" in banner, f"empty-selection banner missing, got: {banner!r}"

            collect_csp(page, "dashboard")

            # ---- Chemistry SPC: dashboard parity/runtime contracts ----
            page.click('#tabs button[data-tab="chem"]')
            page.wait_for_timeout(1800)
            assert page.locator('#chemTopBar .filter-field').count() >= 8, 'Chemistry filter bar did not render'
            assert page.locator('#chemFilterSummaryBar').count() == 1, 'Chemistry selection summary missing'
            assert page.locator('#chemPeriodBanner').count() == 1, 'Chemistry period banner missing'
            assert page.locator('#tab-chem .chem-el-kpi').count() >= 1, 'Chemistry element KPI cards did not render'
            assert page.locator('#tab-chem .chem-el-kpi .sparkline').count() >= 1, 'Chemistry KPI sparkline parity missing'
            # chart-ready is transient, so validate the shared chart classes/structure rather than waiting for a fragile animation state.
            assert page.locator('#chemIChart .chart-svg').count() == 1, 'Chemistry I chart missing shared chart-svg class'
            assert page.locator('#chemMRChart .chart-svg').count() == 1, 'Chemistry MR chart missing shared chart-svg class'
            assert page.locator('#chemHist .chart-svg').count() == 1, 'Chemistry histogram missing shared chart-svg class'
            assert page.locator('#chemIChart .chem-legend').count() == 1, 'Chemistry I chart legend missing'
            assert page.locator('#chemMRChart .chem-legend').count() == 1, 'Chemistry MR chart legend missing'
            assert page.locator('#chemHist .chart-bar').count() >= 1, 'Chemistry histogram is not using the shared chart-bar interaction contract'
            # The I / MR charts show their tooltip through one transparent .chem-hover-layer rect
            # (nearest-heat lookup on pointermove), so drive a real mouse move over that layer;
            # hovering an individual <circle> is intercepted by the layer by design.
            for chart_id in ('chemIChart', 'chemMRChart'):
                shown = False
                for _attempt in range(4):
                    # Scrolling can shift/redraw the chart (sticky bar, resize observer), so scroll,
                    # let it settle, then re-read the layer's box on EVERY attempt.
                    page.evaluate(f"document.querySelector('#{chart_id} .chem-hover-layer').scrollIntoView({{block:'center'}})")
                    page.wait_for_timeout(600)
                    box = page.locator(f'#{chart_id} .chem-hover-layer').bounding_box()
                    assert box, f'{chart_id} hover layer has no box'
                    page.mouse.move(box['x'] + box['width'] * 0.25, box['y'] + box['height'] / 2)
                    page.mouse.move(box['x'] + box['width'] * 0.5, box['y'] + box['height'] / 2, steps=4)
                    page.wait_for_timeout(200)
                    if page.locator('.chart-tooltip.show').count() == 1:
                        shown = True
                        break
                assert shown, f'Chemistry {chart_id} tooltip did not appear'
                tip_txt = page.locator('.chart-tooltip.show .ct-main').inner_text()
                assert tip_txt.strip(), f'Chemistry {chart_id} tooltip is empty'
                page.mouse.move(2, 2)
                page.wait_for_timeout(150)
            month_opts = page.locator('#chemTopBar .filter-field[data-chem-key="month"] .filter-option:not(.all-option)').count()
            if month_opts:
                page.click('#chemTopBar .filter-field[data-chem-key="month"] .filter-trigger')
                page.locator('#chemTopBar .filter-field[data-chem-key="month"] .filter-option:not(.all-option)').first.click()
                page.wait_for_timeout(1200)
                assert page.locator('#chemFilterRecordCount').inner_text().strip() != '-- heats', 'Chemistry Month filter did not refresh the heat count'

            # ---- Chemistry SPC: filter-logic + shared-feature parity with the other tabs ----
            def copts(key):
                return page.eval_on_selector_all(f'#chemTopBar [data-chem-key="{key}"] .filter-option', "e => e.map(x => x.dataset.value)")

            def ctrig(key):
                return page.inner_text(f'#chemTopBar [data-chem-key="{key}"] .filter-trigger span')

            def cpick(key, value):
                page.click(f'#chemTopBar [data-chem-key="{key}"] .filter-trigger')
                page.click(f'#chemTopBar [data-chem-key="{key}"] .filter-option[data-value="{value}"]')
                page.wait_for_timeout(1300)

            page.click('#chemResetAll'); page.wait_for_timeout(1000)
            cpick('spec', 'Test Brass')
            # cascade: like the dashboard, a period you pick narrows the other period lists (no zero-overlap combos can be chosen)
            cpick('quarter', 'Q4')
            assert set(copts('month')) == {'', 'Mar-2026'}, f"Quarter=Q4 should leave only Mar-2026: {copts('month')}"
            assert set(copts('fy')) == {'', 'FY 2025-26'}, f"Quarter=Q4 should leave only FY 2025-26: {copts('fy')}"
            cpick('month', 'Mar-2026')
            # Grade (not the baseline grade) + Quarter + Month = 3 active filters, like a dashboard filter that is not "All"
            assert page.inner_text('#chemActiveBadge') == '3 Active', f"active-filter badge wrong: {page.inner_text('#chemActiveBadge')!r}"
            assert 'filter-active' in (page.get_attribute('#chemTopBar [data-chem-key="spec"]', 'class') or ''), 'changed Grade is not highlighted as active'
            banner = page.inner_text('#chemPeriodBanner')
            assert 'Compared to' in banner and 'Feb-2026' in banner, f'previous-period banner missing/wrong: {banner!r}'
            # dropdown behaviour identical to the dashboard's: search box, one open at a time, click-outside closes
            page.click('#chemTopBar [data-chem-key="week"] .filter-trigger')
            page.fill('#chemTopBar [data-chem-key="week"] .filter-search', 'zzz'); page.wait_for_timeout(150)
            assert page.locator('#chemTopBar [data-chem-key="week"] .filter-empty').count() == 1, 'Chemistry dropdown search has no empty state'
            page.click('#chemTopBar [data-chem-key="param"] .filter-trigger')
            assert page.locator('.filter-control.open').count() == 1, 'opening a 2nd Chemistry dropdown did not close the 1st'
            page.mouse.click(4, 4); page.wait_for_timeout(150)
            assert page.locator('.filter-control.open').count() == 0, 'clicking outside did not close the Chemistry dropdown'
            # switching grade keeps the chosen parameter
            cpick('param', 'ni'); cpick('spec', 'Cu-Ni 80-20')
            assert ctrig('param') == 'Ni%', 'switching grade lost the selected parameter'
            page.click('#chemResetAll'); page.wait_for_timeout(1200)
            assert page.inner_text('#chemActiveBadge') == '0 Active' and ctrig('month') == 'All' and ctrig('quarter') == 'All', 'Reset All did not clear the Chemistry filters'

            # shared features must know about this tab: command palette, number-key shortcut, Reset All command
            cpick('spec', 'Test Brass'); cpick('quarter', 'Q1')
            page.keyboard.press('Control+K'); page.wait_for_timeout(250)
            page.fill('#cmdkInput', 'chemistry'); page.wait_for_timeout(250)
            labels = page.eval_on_selector_all('#cmdkList [data-idx]', "e => e.map(x => x.textContent)")
            assert any('Go to' in l and 'Chemistry SPC' in l for l in labels), f'command palette has no "Go to Chemistry SPC": {labels}'
            page.fill('#cmdkInput', 'Reset All'); page.wait_for_timeout(250); page.keyboard.press('Enter'); page.wait_for_timeout(1300)
            assert ctrig('quarter') == 'All' and page.inner_text('#chemActiveBadge') == '0 Active', 'palette "Reset All Filters" must reset the Chemistry filters while this tab is open'
            page.mouse.click(10, 10); page.keyboard.press('1'); page.wait_for_timeout(900)
            assert not page.evaluate("document.documentElement.classList.contains('chem-mode')"), 'key 1 did not leave the Chemistry tab'
            page.keyboard.press('6'); page.wait_for_timeout(1500)
            assert page.evaluate("document.documentElement.classList.contains('chem-mode')") and page.locator('#tab-chem .chem-el-kpi').count() >= 1, 'key 6 did not open the Chemistry tab'
            # header search result for a Grade only changes dashboard filters -> must land on the Dashboard tab, not change hidden state
            page.click('#globalSearchInput'); page.fill('#globalSearchInput', 'A'); page.wait_for_timeout(1200)
            hit = page.locator('.gsr-item[data-type="grade"]').first
            if hit.count():
                hit.click(); page.wait_for_timeout(1500)
                assert page.evaluate("document.querySelector('.tab-btn.active').dataset.tab") == 'dashboard', 'grade search from the Chemistry tab changed nothing visible'
                page.click('#resetAllBtn'); page.wait_for_timeout(900)
                page.keyboard.press('6'); page.wait_for_timeout(1500)
            # export dialog: honest about scope on this tab
            page.click('#exportMenuBtn'); page.wait_for_timeout(250)
            assert 'Chemistry' in page.inner_text('#exportDialogDesc'), 'export dialog does not say it ignores the Chemistry selection'
            page.keyboard.press('Escape'); page.wait_for_timeout(250)
            # history: Back leaves the tab, Forward returns with data
            page.click('#tabs button:nth-child(1)'); page.wait_for_timeout(700)
            page.click('#tabs button[data-tab="chem"]'); page.wait_for_timeout(1300)
            page.go_back(); page.wait_for_timeout(1300)
            assert not page.evaluate("document.documentElement.classList.contains('chem-mode')"), 'browser Back did not leave Chemistry SPC'
            page.go_forward(); page.wait_for_timeout(1500)
            assert page.locator('#tab-chem .chem-el-kpi').count() >= 1, 'browser Forward did not restore Chemistry SPC'

            # ---- Chemistry SPC: shareable link, Saved Views, Compare Periods, Export dialog ----
            page.click('#chemResetAll'); page.wait_for_timeout(1000)
            cpick('spec', 'Test Brass'); cpick('quarter', 'Q4'); cpick('month', 'Mar-2026'); cpick('param', 'ni')
            share_url = page.url
            for part in ('tab=chem', 'chem_spec=Test+Brass', 'chem_param=ni', 'chem_quarter=Q4', 'chem_month=Mar-2026'):
                assert part in share_url, f'Chemistry selection is not in the address bar ({part}): {share_url}'
            # the link, opened by someone else (fresh browser profile = no remembered selection), restores the same selection
            ctx2 = browser.new_context(viewport={"width": 1440, "height": 900})
            ctx2.route("**/fonts.g*/**", lambda route: route.abort())
            p2 = ctx2.new_page(); watch(p2, "shared-link")
            p2.goto(share_url, wait_until="load")
            p2.click("#introDashboardBtn"); p2.wait_for_timeout(2500)
            def p2trig(key):
                return p2.inner_text(f'#chemTopBar [data-chem-key="{key}"] .filter-trigger span')
            assert p2.evaluate("document.documentElement.classList.contains('chem-mode')"), 'a shared Chemistry link did not open the Chemistry tab'
            assert (p2trig('spec'), p2trig('quarter'), p2trig('month'), p2trig('param')) == (p2.inner_text('#chemTopBar [data-chem-key="spec"] .filter-trigger span'), 'Q4', 'Mar-2026', 'Ni%') and 'Test Brass' in p2trig('spec'), \
                'a shared Chemistry link did not restore grade / period / parameter'
            ctx2.close()
            # Saved Views on this tab (own list, separate from the dashboard's)
            assert page.locator('#chemSaveViewBtn').is_visible(), 'Chemistry tab has no Save Preset button'
            page.click('#chemSaveViewBtn'); page.fill('#chemViewPopSaveName', 'Brass Mar Ni'); page.click('#chemViewPopSaveBtn'); page.wait_for_timeout(300)
            assert 'Brass Mar Ni' in page.eval_on_selector_all('#chemSavedViewSelect option', "e => e.map(x => x.value)"), 'Chemistry preset was not saved'
            assert 'Brass Mar Ni' not in page.evaluate("localStorage.getItem('qdash_saved_views') || ''"), 'Chemistry preset leaked into the dashboard presets'
            page.click('#chemResetAll'); page.wait_for_timeout(1200)
            assert ctrig('month') == 'All', 'Reset All did not clear before loading the preset'
            page.select_option('#chemSavedViewSelect', 'Brass Mar Ni'); page.wait_for_timeout(1500)
            assert (ctrig('quarter'), ctrig('month'), ctrig('param')) == ('Q4', 'Mar-2026', 'Ni%'), 'loading the Chemistry preset did not restore the selection'
            assert 'chem_month=Mar-2026' in page.url, 'loading a Chemistry preset did not update the shareable address'
            page.click('#chemClearViewsBtn'); page.wait_for_timeout(200)
            page.click('#chemViewPopover .view-pop-del'); page.wait_for_timeout(200)
            assert 'Brass Mar Ni' not in page.eval_on_selector_all('#chemSavedViewSelect option', "e => e.map(x => x.value)"), 'deleting a Chemistry preset did not remove it'
            page.click('#chemViewPopCloseBtn')
            # Compare Periods on this tab: two live copies of the Chemistry tab at two different periods
            page.click('#chemResetAll'); page.wait_for_timeout(1000); cpick('spec', 'Test Brass')
            assert page.locator('#chemCompareBtn').is_visible() and not page.locator('#compareModeBtn').is_visible(), 'Chemistry tab must show its own Compare Periods button'
            page.click('#chemCompareBtn'); page.wait_for_timeout(400)
            dims = page.eval_on_selector_all('#compareDimSelect option', "e => e.map(x => x.value)")
            assert dims == ['month', 'week', 'quarter', 'fy', 'spec', 'param'], f'Chemistry compare offers the wrong filters: {dims}'
            page.select_option('#compareDimSelect', 'month'); page.wait_for_timeout(200)
            va, vb = page.input_value('#compareValueA'), page.input_value('#compareValueB')
            assert va and vb and va != vb, f'Chemistry compare did not offer two different months ({va!r}, {vb!r})'
            page.click('#compareGoBtn'); page.wait_for_timeout(3500)
            for fid, want in (('compareFrameA', va), ('compareFrameB', vb)):
                fr = page.frame_locator('#' + fid)
                got = fr.locator('#chemTopBar [data-chem-key="month"] .filter-trigger span').first.inner_text()
                assert got == want, f'{fid}: Chemistry compare pane shows month {got!r}, expected {want!r}'
            assert 'Chemistry' in page.inner_text('#compareViewTitle'), 'Chemistry compare title missing'
            stored_month = page.evaluate("JSON.parse(localStorage.getItem('qdash_chem_sel_v1') || '{}').month || ''")
            assert stored_month == '', f'a Compare Periods pane overwrote the remembered Chemistry selection (month={stored_month!r})'
            page.click('#compareCloseBtn'); page.wait_for_timeout(300)
            # Export dialog: the Chemistry downloads sit in the SAME dialog (only on this tab), next to the disposition reports
            page.click('#exportMenuBtn'); page.wait_for_timeout(300)
            assert page.locator('#exportChemGroup').is_visible() and page.locator('#exportChemGroup .export-dialog-option').count() == 3, 'Chemistry downloads missing from the Export dialog on the Chemistry tab'
            assert page.locator('#exportDialogOptions .export-dialog-option').count() == 4, 'the disposition report options disappeared'
            with page.expect_download() as dl:
                page.click('#exportChemGroup [data-chem-export="cpk"]')
            assert dl.value.suggested_filename.startswith('chemistry_cpk_'), f'unexpected Chemistry export file name: {dl.value.suggested_filename}'
            page.wait_for_timeout(400)
            # dashboard tab: Compare Periods + Export dialog are unchanged (no Chemistry group)
            page.keyboard.press('1'); page.wait_for_timeout(1200)
            page.click('#exportMenuBtn'); page.wait_for_timeout(300)
            assert not page.locator('#exportChemGroup').is_visible(), 'Chemistry downloads must not show in the Export dialog on other tabs'
            page.keyboard.press('Escape'); page.wait_for_timeout(300)
            page.click('#compareModeBtn'); page.wait_for_timeout(300)
            assert 'work_center' in page.eval_on_selector_all('#compareDimSelect option', "e => e.map(x => x.value)"), 'dashboard Compare Periods lost its dimensions'
            page.click('#compareCancelBtn'); page.wait_for_timeout(300)
            page.keyboard.press('6'); page.wait_for_timeout(1200)

            # ---- Admin: login, and every sidebar link must highlight itself ----
            page2 = context.new_page()
            watch(page2, "admin")
            page2.goto(BASE + "/admin", wait_until="load")
            page2.fill("#username", ADMIN_USER)
            page2.fill("#password", ADMIN_PASS)
            page2.click("#loginBtn")
            page2.wait_for_timeout(2000)
            assert page2.eval_on_selector("#loginBox", "el => getComputedStyle(el).display") == "none", \
                "admin login did not succeed"

            mismatches = []
            for link in page2.query_selector_all(".sidebar-link"):
                target = link.get_attribute("data-target")
                link.click()
                page2.wait_for_timeout(1500)
                active = page2.eval_on_selector(".sidebar-link.active", "el => el.dataset.target")
                if active != target:
                    mismatches.append((target, active))
            assert not mismatches, f"admin sidebar highlighted the wrong section: {mismatches}"

            # CSP-safe delegated handlers (#2): users table actions must still work via
            # data-action + addEventListener, not a (now-blocked) onclick="..." attribute.
            # Checked directly on the rendered DOM (not just "does the feature still work")
            # because a stray onclick="..." would be silently inert under this CSP -- the
            # button would still exist and look clickable, it just wouldn't do anything.
            page2.click('.sidebar-link[data-target="usersPanel"]')
            page2.wait_for_timeout(1200)
            reset_btn = page2.query_selector('#usersBody [data-action="reset-user"]')
            assert reset_btn is not None, "reset-user delegated button missing from users table"
            inline_handlers = page2.eval_on_selector_all(
                "#usersBody [onclick]", "els => els.length")
            assert inline_handlers == 0, (
                f"found {inline_handlers} onclick=\"...\" attribute(s) in the users table -- "
                "these are silently blocked by the script-src CSP (no 'unsafe-inline'); "
                "use data-action + addEventListener delegation instead")

            collect_csp(page2, "admin")
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=3)
        except subprocess.TimeoutExpired:
            proc.kill()
        shutil.rmtree(tmpdir, ignore_errors=True)

    assert not violations, "Content-Security-Policy violation(s) at runtime: " + "; ".join(violations)
    assert not page_errors, "Uncaught JS error(s): " + "; ".join(page_errors)
    assert not console_errors, "Browser console error(s): " + "; ".join(console_errors)
    print("BROWSER REGRESSION PASS — dashboard load, all tabs, cascading filters, "
          "empty-selection banner, Chemistry SPC filter cascade + shared-feature parity, admin nav highlighting, CSP-safe admin actions; "
          "zero CSP violations / console errors / uncaught JS errors.")


if __name__ == "__main__":
    run()
