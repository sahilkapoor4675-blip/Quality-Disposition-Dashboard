#!/usr/bin/env python3
"""Chemistry SPC checks: SPC maths vs hand calculation, import validation rules, spec parsing/matching,
and the HTTP flow (preview -> confirm -> SPC view -> heat join -> backup round-trip) on a temporary database.
Run: python tests/test_chem_spc.py"""
import datetime as dt, io, json, os, sys, tempfile, threading, http.client, statistics as st
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import openpyxl
import chem_spc as c

ERR = []
def check(cond, msg):
    if not cond: ERR.append(msg)

# ---------------------------------------------------------------- SPC maths
v = [10, 12, 11, 13, 12, 14, 11, 10, 12, 13]
ch = c.imr(v); mr = [abs(v[i] - v[i - 1]) for i in range(1, len(v))]
check(abs(ch["mrbar"] - sum(mr) / len(mr)) < 1e-12, "MRbar")
check(abs(ch["ucl"] - (sum(v) / len(v) + 3 * (sum(mr) / len(mr)) / 1.128)) < 1e-9, "I-chart UCL = mean + 3*MRbar/d2")
check(abs(ch["mr_ucl"] - 3.267 * ch["mrbar"]) < 1e-9, "MR UCL = D4*MRbar")
cap = c.capability(v, 8, 16, ch["sigma_within"]); m, s, sw = sum(v) / 10, st.stdev(v), ch["sigma_within"]
check(abs(cap["cp"] - 8 / (6 * sw)) < 1e-9 and abs(cap["cpk"] - min((16 - m), (m - 8)) / (3 * sw)) < 1e-9, "Cp/Cpk")
check(abs(cap["pp"] - 8 / (6 * s)) < 1e-9 and abs(cap["ppk"] - min((16 - m), (m - 8)) / (3 * s)) < 1e-9, "Pp/Ppk")
one = c.capability([1, 2, 1.5, 3], None, 5, 1.0); check(one["cp"] is None and one["cpk"] is not None, "one-sided spec has Cpk only")
check(not hasattr(c, "western_electric"), "Western Electric rules are gone")
flat = c.analyse_param([{"value": 0.0}] * 10, {"al": [0, .5]}, "al"); check(any("identical" in w for w in flat["warnings"]), "flat data warning")
check(c.eff_limits({"pb": [0, .005]}, "pb") == (None, .005), "LSL 0 ignored")

# ---------------------------------------------------------------- import validation
def rows(*recs):
    out = []
    for i, d in enumerate(recs, start=2):
        base = {"heat_no": "NBS1", "cast_date": "01.05.2026", "alloy": "NBS", "denomination": "5RS", "analyst": "SUNIL", "cu": 75.0, "ni": 5.0, "zn": 19.95, "mn": .01, "fe": .02, "pb": .003, "_sheet": "S", "_row": i}
        base.update(d); out.append(base)
    return out
res = c.validate_rows(rows({}, {"heat_no": "NBS1"}))
check(res["duplicates"] == 1 and res["errors"] == 0 and len(res["records"]) == 1, "identical duplicate heat is skipped, not an error")
res = c.validate_rows(rows({}, {"heat_no": "NBS1", "cu": 74.5}))
check(res["errors"] == 2 and not res["records"], "conflicting duplicate heat rejects every copy")
# cast date: read day-first, stored as ISO, only feeds the period filters; a bad date never rejects a heat
P = c.parse_cast_date
check(P("02.04.2026") == ("2026-04-02", "") and P("04.04.26") == ("2026-04-04", "") and P("2026-04-02") == ("2026-04-02", ""), "dd.mm.yyyy, dd.mm.yy and ISO dates are read day-first")
check(P(dt.datetime(2026, 4, 2, 10, 30)) == ("2026-04-02", "") and P(46114)[0] == "2026-04-02", "Excel datetime and serial dates are read")
check(P(2505.2025) == ("2025-05-25", "guessed"), "a date typed as the number 2505.2025 is reconstructed and flagged")
check(P("27.05.206")[0] == "" and P("not a date")[0] == "" and P("31.02.2026")[0] == "" and P("")[1] == "blank" and P(None)[1] == "blank", "unreadable / blank dates give no date")
check(c.period_fields("2026-04-02") == ("Apr-2026", "Wk of 30-Mar-26", "Q1", "FY 2026-27") and c.period_fields("2027-02-10")[2:] == ("Q4", "FY 2026-27") and c.period_fields("2026-03-31")[3] == "FY 2025-26", "Month/Week/Quarter/FY labels match the main dashboard (Apr-Mar financial year)")
for bad_date, code in (("01.01.2099", "future_date"), ("not a date", "bad_date"), (2505.2025, "date_guess")):
    res = c.validate_rows(rows({"heat_no": "NBS2", "cast_date": bad_date}))
    check(res["errors"] == 0 and len(res["records"]) == 1 and res["issue_counts"].get(code) == 1, f"date value {bad_date!r} only warns ({code}); the heat is still imported")
res = c.validate_rows(rows({"heat_no": "NBS2", "cast_date": "02.04.26"})); check(res["records"][0]["cast_date"] == "2026-04-02" and res["warnings"] == 0, "a good date is stored as ISO with no warning")
res = c.validate_rows(rows({"heat_no": "NBS2", "cast_date": None})); check(res["records"][0]["cast_date"] == "" and res["warnings"] == 0, "a blank date is only informational")
check(c._map_header(["Date", "Coil No.", "Cu%"]) == {0: "cast_date", 1: "heat_no", 2: "cu"}, "the Date column of the chemistry file is mapped")
recs_p = [{"heat_no": f"H{i}", "cast_date": d, "cu": 75 + i * .01} for i, d in enumerate(["2026-04-02", "2026-04-09", "2026-05-03", "2026-07-01", ""])]
check([r["heat_no"] for r in c.filter_by_period(recs_p, {"month": "Apr-2026"})] == ["H0", "H1"] and [r["heat_no"] for r in c.filter_by_period(recs_p, {"quarter": "Q1", "fy": "FY 2026-27"})] == ["H0", "H1", "H2"] and len(c.filter_by_period(recs_p, {"week": "Wk of 30-Mar-26"})) == 1 and len(c.filter_by_period(recs_p, {})) == 5 and len(c.filter_by_period(recs_p, {"month": "All"})) == 5, "period filters (undated heats drop out only when a filter is on)")
po = c.period_options(recs_p); check(po["months"][0] == "Jul-2026" and po["undated"] == 1 and po["quarters"] == ["Q1", "Q2"] and po["fys"] == ["FY 2026-27"], "period options newest first, undated counted")
res = c.validate_rows(rows({"heat_no": "NBS4", "cu": 175.0, "total": None})); check(res["issue_counts"].get("out_of_range") == 1, "% above 100")
res = c.validate_rows(rows({"heat_no": "NBS5", "ni": -1})); check(res["issue_counts"].get("out_of_range") == 1, "negative %")
res = c.validate_rows(rows({"heat_no": "NBS6", "cu": "abc"})); check(res["issue_counts"].get("bad_number") == 1, "text in a number cell")
res = c.validate_rows(rows({"heat_no": "NBS7", "total": 90.0})); check(res["issue_counts"].get("total_mismatch") == 1 and res["errors"] == 1, "Total% far from sum -> error")
res = c.validate_rows(rows({"heat_no": "NBS8", "total": 100.1})); check(res["errors"] == 0 and res["warnings"] >= 1, "Total% slightly off -> warning only")
res = c.validate_rows(rows({"heat_no": ""})); check(res["issue_counts"].get("missing_heat") == 1, "blank heat")
existing = {"NBS1": {**c.validate_rows(rows({}))["records"][0]}}
res = c.validate_rows(rows({}), existing); check(res["unchanged"] == 1 and not res["records"], "re-import identical = unchanged")
res = c.validate_rows(rows({"cu": 75.2}), existing); check(res["updated"] == 1 and res["updated_details"][0]["changes"][0]["field"] == "Cu%", "changed value = update with diff")
names = rows(*[{"heat_no": f"NBS{i}", "analyst": "JANMEJAY"} for i in range(20, 32)], {"heat_no": "NBS40", "analyst": "JANMJAY"})
res = c.validate_rows(names); check(res["issue_counts"].get("analyst_typo") == 1, "analyst name typo")

# re-import semantics: numbers are overwritten, but a blank analyst/sheet never erases what is stored
first = c.validate_rows(rows({}))["records"][0]; first.pop("_status", None)
res = c.validate_rows(rows({"analyst": None, "cu": 75.2, "zn": 19.75}), {"NBS1": first})
rec = res["records"][0]
check(rec["_status"] == "update" and rec["cu"] == 75.2 and rec["analyst"] == "SUNIL", "re-import overwrites numbers but keeps the stored analyst when the file leaves it blank")
res = c.validate_rows(rows({"cast_date": "02.05.2026"}), {"NBS1": first})
check(res["updated"] == 1 and res["updated_details"][0]["changes"][0]["field"] == "cast_date", "a corrected date in the file is an update")
res = c.validate_rows(rows({"cast_date": None}), {"NBS1": first})
check(res["unchanged"] == 1 and not res["records"], "a blank date never erases the stored one")
res = c.validate_rows(rows({"_sheet": "Other sheet"}), {"NBS1": {**first, "sheet": "S"}})
check(res["updated"] == 1, "moving a heat to another sheet is an update (spec assignment depends on it)")

# ---------------------------------------------------------------- heat-number order, 3-decimal precision, main elements (V66.1 chemistry follow-up)
hs = ["NBS1000", "NBS999", "NSC12", "NBS20", "NSC5", "PBC1"]
check([r["heat_no"] for r in sorted(({"heat_no": h, "cast_date": "2030-01-0%d" % (i + 1)} for i, h in enumerate(hs)), key=c.order_key)] == ["NBS20", "NBS999", "NBS1000", "NSC5", "NSC12", "PBC1"], "order = prefix then NUMERIC heat sequence, whatever the dates say")
v = c.build_spc_view([{"heat_no": h, "cu": 75 + i * .001} for i, h in enumerate(["NBS10", "NBS2", "NBS1", "NBS3"])], None, "cu", {}, last_n=2)
check([p["heat_no"] for p in v["series"]] == ["NBS3", "NBS10"], "last N = highest heat numbers")
check(c.CHEM_DIGITS == 3, "chemistry precision is 3 decimals")
lim = {"total": [99.9, 100.0], "cu": [74, 76]}
check(not c.spec_violations({"total": 100.0004, "cu": 75}, lim) and c.spec_violations({"total": 100.004, "cu": 75}, lim)[0]["param"] == "total", "Total% is compared at 3 decimals (100.004 is out, 100.0004 is noise)")
check(c.spec_violations({"cu": 76.0004}, {"cu": [74, 76]}) and not c.spec_violations({"cu": 76.0}, {"cu": [74, 76]}), "element values are compared exactly, no rounding below 3 decimals")
res = c.validate_rows(rows({"heat_no": "NBS90", "cu": 75.123, "ni": 5.001, "zn": 19.876, "pb": 0.003, "mn": 0, "fe": 0, "total": 100.003}))
r0 = res["records"][0]; check(res["errors"] == 0 and (r0["cu"], r0["ni"], r0["zn"], r0["pb"]) == (75.123, 5.001, 19.876, 0.003), "3-decimal values are stored exactly as reported")
check([c.cpk_rating(x) for x in (None, 2.0, 1.67, 1.5, 1.33, 1.2, 1.0, 0.99)] == [None, "excellent", "excellent", "capable", "capable", "marginal", "marginal", "poor"], "Cpk rating bands 1.67 / 1.33 / 1.00")
check(c.main_elements({"cu": [74, 76], "zn": [19, 21], "pb": [0, .04], "fe": [None, .1]}, {"cu": 75, "zn": 20, "pb": .003}) == ["cu", "zn"], "main elements = Cu + elements with a real LSL; impurities excluded")
check(c.main_elements({"cu": [74, 76], "ni": [4.5, 5.5], "zn": [18, 21]}, {"cu": 75, "ni": 5, "zn": 20}) == ["cu", "zn", "ni"], "main elements ordered Cu first, then by mean")
check(c.main_elements({"pb": [0, .04]}, {"cu": 99.9, "pb": .003, "zn": .2, "ni": 1.5}) == ["cu", "ni"], "no alloying LSL in the spec -> elements averaging >= 1 % are main")
check(c.main_elements({}, {}) == ["cu"], "Cu is always a main element")

# ---------------------------------------------------------------- spec parsing / matching
def xlsx(header, body):
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(header)
    for b in body: ws.append(b)
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()
hdr = ["Alloy", "Grade Descriptions", "Cu% (LSL)", "Cu% (USL)", "Pb% (LSL)", "Pb% (USL)"]
specs, issues = c.parse_spec_file("s.xlsx", xlsx(hdr, [["NBS", "Brass (5rs.)", 74, 76, 0, .04], ["NBS", "Brass (10rs. & 20rs.)", 74, 76, 0, .005], ["X", "Bad", 80, 70, None, None], ["Y", "Brass (5rs.)", 1, 2, None, None]]))
check(len(specs) == 2 and {i["code"] for i in issues} >= {"lsl_gt_usl", "dup_description"}, "spec file: LSL>USL and duplicate description are reported")
sp = c.resolve_spec(specs, "NBS", "anything", "5RS"); check(sp and sp["description"] == "Brass (5rs.)", "alloy with two specs resolved by denomination")
sp = c.resolve_spec(specs, "NBS", "Brass (10rs. & 20rs.)", "5RS"); check(sp and "10rs" in sp["description"], "sheet name wins over denomination")
check(c.resolve_spec(specs, "NBS", "", "") is None, "ambiguous spec is not guessed")
def xlsx2(header, std_rows, aim_rows, aim_title="AIM"):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Standard"; ws.append(header)
    for b in std_rows: ws.append(b)
    w2 = wb.create_sheet(aim_title); w2.append(header)
    for b in aim_rows: w2.append(b)
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()
sp2, is2 = c.parse_spec_file("s.xlsx", xlsx2(hdr, [["NBS", "Brass (5rs.)", 74, 76, 0, .04], ["NBS", "Cu-Ni", 86, 88, None, None]], [["NBS", "Brass (5rs.)", 74.2, 75.8, 0, .01], ["X", "Ghost", 1, 2, None, None]]))
check(len(sp2) == 2 and sp2[0]["limits"]["cu"] == [74, 76] and sp2[0]["aim"]["cu"] == [74.2, 75.8] and sp2[0]["aim"]["pb"] == [0, .01] and sp2[1]["aim"] == {}, "Standard sheet + AIM sheet are merged per grade")
check({i["code"] for i in is2} >= {"aim_without_standard", "no_aim"}, "AIM row without a Standard row / Standard row without AIM are reported")
sp3, is3 = c.parse_spec_file("s.xlsx", xlsx2(hdr, [["NBS", "Brass (5rs.)", 74, 76, 0, .04]], [["NBS", "Brass (5rs.)", 73, 77, 0, .01]]))
check(any(i["code"] == "aim_outside_standard" for i in is3), "AIM wider than Standard is flagged")
sp4, is4 = c.parse_spec_file("s.xlsx", xlsx2(hdr, [["NBS", "Brass (5rs.)", 74, 76, 0, .04], ["NBS", "Cu-Ni", 86, 88, None, None]], [["NBS", "Brass (5rs.) AIM", 74.2, 75.8, 0, .01], ["X", "Ghost", 1, 2, None, None]]))
check(sp4[0]["aim"].get("cu") == [74.2, 75.8] and any(i["code"] == "aim_loose_match" for i in is4) and any(i["code"] == "aim_without_standard" for i in is4) and sp4[1]["aim"] == {}, "an AIM grade name that only differs by a suffix is matched to its single Standard grade (reported); unrelated names are still ignored")
check(c.parse_spec_file("s.xlsx", xlsx(hdr, [["NBS", "Brass", 74, 76, 0, .04]]))[0][0]["aim"] == {}, "a single-sheet spec file gives Standard limits only")
check(c.validate_aim_payload({"cu": ["74.2", "75.8"]}) == ({"cu": [74.2, 75.8]}, None) and c.validate_aim_payload({"cu": [2, 1]})[1] and c.validate_aim_payload({"zz": [1, 2]})[1] and c.validate_aim_payload(None) == ({}, None), "manual AIM limits validated")
check(c.validate_spec_payload("A", "d", {"cu": [1, 0]})[1], "manual spec LSL>USL rejected")
check(c.validate_spec_payload("A", "d", {"zz": [1, 2]})[1], "manual spec unknown parameter rejected")

# ---------------------------------------------------------------- regressions found in review (V66.1 hardening)
check(c.parse_number("75.2 %") == 75.2 and c.parse_number("0.003%") == 0.003, "a % sign typed into a number cell is accepted")
res = c.validate_rows(rows({"heat_no": "NBS70", "cu": "75.2 %", "total": None})); check(res["errors"] == 0 and res["records"][0]["cu"] == 75.2, "text '75.2 %' imports as 75.2")
# Excel truncates sheet names at 31 chars: a truncated '... (10rs. & 20r' must not raise a false denomination warning for 20 Rs heats
long_sheet = "NI-Brass (Ni - 05) (10rs. & 20r"
res = c.validate_rows(rows({"heat_no": "NBS71", "denomination": "20RS", "_sheet": long_sheet})); check(res["issue_counts"].get("denom_sheet") is None, "truncated sheet name does not trigger a false denomination warning")
res = c.validate_rows(rows({"heat_no": "NBS72", "denomination": "20RS", "_sheet": "Brass (5rs.)"})); check(res["issue_counts"].get("denom_sheet") == 1, "a genuine denomination/sheet mismatch is still reported")
# re-import from a file without sheet/alloy/denomination columns must keep resolving to the stored heat's spec
sp_list = [{"alloy": "NBS", "description": "S", "limits": {"cu": [74, 76]}}]
stored = c.validate_rows(rows({"heat_no": "NBS80", "_sheet": "S"}), None, sp_list)["records"][0]
res = c.validate_rows(rows({"heat_no": "NBS80", "_sheet": "", "alloy": "", "denomination": "", "cu": 77.0, "total": None}), {"NBS80": stored}, sp_list)
check(res["records"] and res["records"][0]["_spec"] == "S" and not res["unresolved_sheets"] and res["oos_heats"] == 1, "re-import without sheet/alloy keeps the stored spec (no false 'no spec' warning, OOS still counted)")
v = c.build_spc_view([{"heat_no": f"A{i}", "cast_date": d, "cu": 75 + i * .01} for i, d in enumerate(["2026-04-02", "2026-05-02", "2026-05-09"])], {"description": "S", "alloy": "N", "limits": {"cu": [74, 76]}, "aim": {"cu": [74.2, 75.8]}}, "cu", {}, period={"month": "May-2026"})
check(v["n"] == 2 and (v["aim_lsl"], v["aim_usl"], v["lsl"], v["usl"]) == (74.2, 75.8, 74, 76) and all("rules" not in p for p in v["series"]) and "ooc_points" not in v["summary"] and "out_of_control" not in v["summary"]["compare"] and v["series"][0]["cast_date"] == "2026-05-02", "SPC view: period filter, Standard + Aim limits, no Western Electric fields")
check(v["capability"]["sigma_within"] is not None and v["capability"]["sigma_overall"] is not None, "capability carries both standard deviations")
# a group with no limits must not claim its heats are 'in spec'
v = c.build_spc_view([{"heat_no": "X1", "cast_date": "2026-05-01", "cu": 75.0}, {"heat_no": "X2", "cast_date": "2026-05-02", "cu": 75.1}], None, "cu", {"X1": c.summarize_disposition([{"output_weight": 1, "quality_decision": "ACCEPT"}])})
check(v["summary"]["has_limits"] is False and v["summary"]["compare"]["in_spec"]["heats"] == 0, "no limits -> no heat is reported as 'in spec'")
v = c.build_spc_view([{"heat_no": "X1", "cast_date": "2026-05-01", "cu": 75.0}, {"heat_no": "X2", "cast_date": "2026-05-02", "cu": 75.1}], {"description": "S", "alloy": "N", "limits": {"cu": [74, 76]}}, "cu", {"X1": c.summarize_disposition([{"output_weight": 1, "quality_decision": "ACCEPT"}])})
check(v["summary"]["has_limits"] is True and v["summary"]["compare"]["in_spec"]["heats"] == 1, "with limits, in-spec heats are still counted")

# ---------------------------------------------------------------- HTTP flow
def http_flow():
    import shutil, urllib.request, urllib.error, re
    td = tempfile.mkdtemp(prefix="qdash_chem_")
    shutil.copy(ROOT / "quality.db", Path(td) / "quality.db")
    os.environ.update(DB_PATH=str(Path(td) / "quality.db"), APP_VERSION="test", BACKUP_DIR=str(Path(td) / "b"), ADMIN_USERNAME="admin", ADMIN_PASSWORD="Strong-Admin-1234!")
    os.environ.pop("RENDER", None)
    import server
    server._ensure_admin_schema(); server.ensure_fast_indexes(); server.STARTUP_READY = True
    from http.server import ThreadingHTTPServer
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), server.Handler); port = httpd.server_address[1]
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    try:
        login = urllib.request.Request(f"http://127.0.0.1:{port}/api/login", data=json.dumps({"username": "admin", "password": "Strong-Admin-1234!"}).encode(), headers={"Content-Type": "application/json"}, method="POST")
        with urllib.request.urlopen(login) as rr: cookie = "; ".join(h.split(";")[0] for h in rr.headers.get_all("Set-Cookie"))
        csrf = re.search(r"qdash_csrf=([^;]+)", cookie).group(1)
        def call(method, path, body=None, files=None, auth=True):
            h = {"Cookie": cookie, "X-CSRF-Token": csrf} if auth else {}
            data = None
            if files:
                b = "----t" + "x" * 12; parts = []
                for name, (fn, content) in files.items():
                    parts.append(f'--{b}\r\nContent-Disposition: form-data; name="{name}"; filename="{fn}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode() + content + b"\r\n")
                data = b"".join(parts) + f"--{b}--\r\n".encode(); h["Content-Type"] = f"multipart/form-data; boundary={b}"
            elif body is not None:
                data = json.dumps(body).encode(); h["Content-Type"] = "application/json"
            req = urllib.request.Request(f"http://127.0.0.1:{port}{path}", data=data, headers=h, method=method)
            try:
                with urllib.request.urlopen(req) as x: return x.status, json.loads(x.read().decode())
            except urllib.error.HTTPError as e: return e.code, json.loads(e.read().decode() or "{}")
        # specs
        std = xlsx2(["Alloy", "Grade Descriptions", "Cu% (LSL)", "Cu% (USL)", "Zn% (LSL)", "Zn% (USL)", "Pb% (LSL)", "Pb% (USL)"], [["NBS", "Test Brass", 74, 76, 19, 21, 0, .04]], [["NBS", "Test Brass", 74.2, 75.8, 19.2, 20.8, 0, .01]])
        s, d = call("POST", "/api/admin/chem_specs_preview", files={"file": ("Standard.xlsx", std)}); check(s == 200 and d["new"] == 1, f"spec preview {s} {d}")
        s, d = call("POST", "/api/admin/chem_specs_confirm", {"preview_id": d["preview_id"]}); check(s == 200 and d["inserted"] == 1, f"spec confirm {s} {d}")
        # chemistry: heats, one matching a real disposition heat (NBS6348), one out of spec, one bad row
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Test Brass"
        ws.append(["Date", "Coil No.", "Alloy", "Denomination", "Cu%", "Ni%", "Zn%", "Pb%", "Total%", "Name"])
        for i in range(30):
            hn = "NBS6348" if i == 0 else f"NBS9{i:03d}"
            cu = 75.0 + (i % 5 - 2) * 0.05; zn = 20 - (i % 5 - 2) * 0.05
            ws.append([f"{1 + i % 27:02d}.05.2026", hn, "NBS", "5RS", cu, 5.0, zn, .003, cu + 5.0 + zn + .003, "SUNIL"])
        ws.append(["10.05.2026", "NBS9900", "NBS", "5RS", 77.0, 5.0, 17.99, .003, 100.0, "SUNIL"])      # Cu above USL, Zn below LSL
        ws.append(["10.05.2026", "NBS9901", "NBS", "5RS", 75.0, 5.0, 20.0, .003, 100.003, "SUNIL", ]); ws.append(["10.05.2026", "NBS9902", "NBS", "5RS", "oops", 5.0, 20.0, .003, 100.0, "SUNIL"])
        bio = io.BytesIO(); wb.save(bio)
        s, d = call("POST", "/api/admin/chem_import_preview", files={"file": ("chem.xlsx", bio.getvalue())})
        check(s == 200 and d["new"] == 32 and d["errors"] == 1 and d["oos_heats"] >= 1 and d["dated"] == 32 and d["undated"] == 0 and d["date_from"] == "2026-05-01" and all(r.get("cast_date") for r in d["sample"]), f"chem preview {s} { {k: d.get(k) for k in ('new', 'errors', 'oos_heats')} }")
        s2, _ = call("POST", "/api/admin/chem_import_confirm", {"preview_id": "nope"}); check(s2 == 400, "unknown preview token is refused")
        s, d2 = call("POST", "/api/admin/chem_import_confirm", {"preview_id": d["preview_id"]}); check(s == 200 and d2["inserted"] == 32, f"chem confirm {s} {d2}")
        s, d3 = call("POST", "/api/admin/chem_import_confirm", {"preview_id": d["preview_id"]}); check(s == 400, "a preview token cannot be confirmed twice")
        s, d = call("POST", "/api/admin/chem_import_preview", files={"file": ("chem.xlsx", bio.getvalue())}); check(d["new"] == 0 and d["unchanged"] == 32, "re-preview shows every heat unchanged")
        s, _ = call("POST", "/api/admin/chem_import_preview", files={"file": ("chem.xlsx", bio.getvalue())}, auth=False); check(s in (401, 403), f"chem import requires admin login ({s})")
        # public SPC + heat join
        s, m = call("GET", "/api/chem/meta", auth=False); check(s == 200 and m["specs"][0]["heats"] == 32 and m["specs"][0]["with_disposition"] >= 1 and m["specs"][0]["aim"]["cu"] == [74.2, 75.8] and m["specs"][0]["periods"]["months"] == ["May-2026"], f"meta {s} (Aim limits + period options)")
        q = "/api/chem/spc?spec=Test%20Brass&param=cu"
        s, v = call("GET", q, auth=False)
        check(s == 200 and v["n"] == 32 and v["oos_total"] >= 1 and v["imr"]["ucl"] > v["imr"]["cl"] > v["imr"]["lcl"], f"spc view {s}")
        check(any(o["heat_no"] == "NBS9900" and {x["key"] for x in o["violations"]} == {"cu", "zn"} for o in v["oos_heats"]), "out-of-spec list shows both breached parameters")
        check(any(p["heat_no"] == "NBS6348" and p["disp"] and p["disp"]["coils"] >= 1 for p in v["series"]), "point carries disposition summary joined on heat_no")
        check([p["heat_no"] for p in v["series"]] == sorted((p["heat_no"] for p in v["series"]), key=lambda h: (c.heat_prefix(h), c.heat_seq(h))) and v["series"][0]["heat_no"] == "NBS6348" and all(p["cast_date"] for p in v["series"]), "series is in heat-number order (dates are carried but never order it)")
        ov = {o["param"]: o for o in v["overview"]}
        check([o["param"] for o in v["overview"]][:2] == ["cu", "zn"] and ov["cu"]["main"] and ov["zn"]["main"] and not ov["pb"]["main"], f"overview lists main elements (Cu, Zn) first: {[o['param'] for o in v['overview']]}")
        check(all(o["rating"] == c.cpk_rating(o["cpk"]) for o in v["overview"]) and v["cpk_bands"] == {"excellent": 1.67, "capable": 1.33, "marginal": 1.0}, "overview carries the Cpk rating and the bands")
        s, v2 = call("GET", q + "&date_from=2026-05-10&date_to=2026-05-10", auth=False); check(s == 200 and v2["n"] == 32, f"old date_from/date_to parameters are ignored, all heats stay ({v2.get('n')})")
        s, v2 = call("GET", q + "&month=May-2026&quarter=Q1&fy=FY%202026-27", auth=False); check(s == 200 and v2["n"] == 32 and v2["aim_lsl"] == 74.2 and v2["aim_usl"] == 75.8, "period filter + Aim limits in the SPC view")
        s, v2 = call("GET", q + "&month=Jun-2026", auth=False); check(s == 200 and v2["n"] == 0 and v2["overview"] == [], "a period with no heats gives an empty view, not an error")
        check(all(o["sigma_within"] > 0 and o["sigma_overall"] > 0 for o in v["overview"] if o["param"] in ("cu", "zn")), "overview carries both standard deviations")
        s, v3 = call("GET", q + "&last_n=10", auth=False); check(s == 200 and v3["n"] == 10, "last N")
        s, _ = call("GET", "/api/chem/spc?spec=Nope&param=cu", auth=False); check(s == 400, "unknown spec -> 400")
        s, _ = call("GET", "/api/chem/spc?spec=Test%20Brass&param=zz", auth=False); check(s == 400, "unknown param -> 400")
        s, h = call("GET", "/api/chem/heat?heat_no=nbs6348", auth=False); check(s == 200 and h["found"] and h["coils"] and h["summary"]["coils"] == len(h["coils"]), "heat dialog data (case-insensitive)")
        # manual spec edit
        s, e = call("POST", "/api/admin/chem_spec", {"description": "Manual", "alloy": "M", "limits": {"cu": [1, 0]}}); check(s == 400, "manual spec with LSL>USL refused")
        s, e = call("POST", "/api/admin/chem_spec", {"description": "Manual", "alloy": "M", "limits": {"cu": ["1", "2"]}}); check(s == 200, f"manual spec saved {s} {e}")
        s, e = call("POST", "/api/admin/chem_spec", {"description": "manual", "alloy": "M", "limits": {"cu": [1, 2]}}); check(s == 400, "duplicate spec name (case-insensitive) refused")
        # backup round trip keeps chemistry
        snap = server._backup_snapshot_transaction("chem-test")
        check({"chem_heats", "chem_specs", "chem_import_history"} <= set(snap["tables"]) and len(snap["tables"]["chem_heats"]) == 32, "chemistry tables are in backups")
        conn = server.get_conn(); conn.execute("DELETE FROM chem_heats"); conn.commit(); conn.close()
        server._restore_backup_data(snap); check(server.compute_chem_meta()["total_heats"] == 32, "restore brings chemistry back")
        # a recovery point made BEFORE the chemistry tables existed must still restore, leaving chemistry untouched
        import copy
        older = copy.deepcopy(server._backup_snapshot_transaction("old-version"))
        older["persistent_tables"] = [t for t in older["persistent_tables"] if not t.startswith("chem_")]
        for k in ("tables", "table_schema", "counts"): older[k] = {t: v for t, v in older[k].items() if not t.startswith("chem_")}
        older["schema_fingerprint"] = server._schema_fingerprint(older["table_schema"]); older["integrity_sha256"] = server._backup_integrity_sha256(older)
        try:
            server._restore_backup_data(older); check(server.compute_chem_meta()["total_heats"] == 32, "pre-chemistry backup restores and leaves chemistry data untouched")
        except Exception as ex: check(False, f"pre-chemistry backup could not be restored: {ex}")
    finally:
        httpd.shutdown(); httpd.server_close()
http_flow()

# ---------------------------------------------------------------- source contracts
idx = (ROOT / "index.html").read_text(encoding="utf-8"); js = "".join(p.read_text(encoding="utf-8") for p in sorted((ROOT / "src/js").glob("*.js")))
check('data-tab="chem"' in idx and 'id="tab-chem"' in idx, "dashboard tab + panel present")
check(("chem: loadChemSpc" in js and "'chem'" in js) or "TAB_LOADERS.chem = loadChemSpc" in js, "tab registered")
adm = (ROOT / "admin.html").read_text(encoding="utf-8")
check('id="chemImportPanel"' in adm and 'id="chemSpecPanel"' in adm and ' onclick="' not in adm, "admin panels present, no inline handlers")
check('id="chemMainBox"' in idx and "renderChemMain" in js and 'id="chemKpis"' not in idx and "renderChemKpis" not in js and "URLSearchParams({spec: chemSel.spec, param: chemSel.param, last_n" in js and "month: chemSel.month" in js, "element cards present; the KPI cards are gone; period filters are sent")
check("Western" not in js and "rules" not in js.split("function drawChemI")[1].split("function chemLegend")[0] and "Western Electric" not in idx, "no Western Electric rule anywhere in the Chemistry SPC tab")
check(all(k in js for k in ("Std. Dev.", "chemElemArt", "chemFieldHtml", "chemFindHeat", "chemResetAll", "kpi-card", "Aim LSL", "Std LSL")) and "chemInsIcon" not in js,
      "Dashboard-style filter bar, KPI-style element cards, Std. Dev. cells, Standard + Aim lines present; Insert box removed")
check("data-aim" in adm and "AIM limits" in adm and "Date span" in adm, "admin: AIM limits editor and chemistry date summary present")
css = (ROOT / "src/css/16-chem-spc.css").read_text(encoding="utf-8")
check("#0f2a4a" not in js and 'stroke="var(--text)"' in js, "histogram normal curve follows the theme (no fixed dark navy)")
check('html[data-theme="dark"] .chem-great' in css, "dark-theme colour for the Excellent Cpk rating")
check("chemExportCpk" in js and "chemExportHeats" in js and 'id="chemCpkCsv"' in js and 'id="chemHeatsCsv"' in js, "Cpk-table and heat-data CSV exports present")
check("TAB_LOADERS.chem = loadChemSpc" in js and "#tabs .tab-btn" in css, "tab loader registered from the chemistry piece; tab bar stays on one row")
check("row" in js.split("function drawChemHist")[1] and "Mean" in js.split("function drawChemHist")[1], "histogram labels use separate rows for Standard, Aim and Mean")

# ---------------------------------------------------------------- period filters must get data (dates with a time part, any "…Date…" header, cascading lists)
import chem_spc as _cs
check(all(_cs.parse_cast_date(v)[0] == "2026-04-02" for v in ("02.04.2026 10:30:00", "2026-04-02T10:30:00", "2-Apr-26", "02/04/2026", "Apr 2, 2026", "20260402")),
      "dates typed with a time part / other common formats are read (they used to leave every heat undated -> empty Month/Week/Quarter/FY)")
check(_cs.parse_cast_date("2505.2025") == ("2025-05-25", "guessed"), "a day.month.year typed as text 2505.2025 is read like the numeric case")
check(all(_cs._is_date_header(_cs._hnorm(h)) for h in ("Date", "Cast Date", "Date of Analysis", "Analysis Date", "Sample Date", "Date/Time")) and not _cs._is_date_header(_cs._hnorm("Updated At")),
      "any header that says 'date' is the cast date column")
_recs = [{"heat_no": f"NB{i}", "cast_date": d} for i, d in enumerate(["2026-01-05", "2026-01-20", "2026-02-10", "2026-04-15", "2026-07-01"])]
_c = _cs.period_options_cascade(_recs, {"month": "Jan-2026", "quarter": "Q4", "fy": "FY 2025-26"})
check(_c["months"] == ["Feb-2026", "Jan-2026"] and _c["quarters"] == ["Q4"] and _c["fys"] == ["FY 2025-26"] and all(w.startswith("Wk of") for w in _c["weeks"]),
      "period dropdown lists cascade: each list follows the OTHER three selections, never its own")

if ERR:
    print("CHEM SPC FAIL"); [print(" -", e) for e in ERR]; sys.exit(1)
print("CHEM SPC PASS — SPC maths, import validation, spec matching, HTTP flow, heat join, backup round-trip.")
