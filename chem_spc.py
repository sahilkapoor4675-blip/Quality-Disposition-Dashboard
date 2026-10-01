"""Cast-chemistry import + Statistical Process Control (SPC) logic.

Pure functions only (no HTTP, no database) so everything here is unit-testable.
server.py owns the tables, endpoints, locks and backups and calls into this module.

Concepts
--------
* One *heat* (the "Coil No." column of the chemistry workbook == ``heat_no`` in the
  disposition table) is one SPC data point.
* Heats are ordered by heat number ONLY (letters prefix, then the numeric sequence). The cast date of a heat
  (the "Date" column of the chemistry workbook) is read and stored ONLY to drive the Month / Week / Quarter /
  Financial-year filters; it never orders the charts and an unreadable date never rejects a heat (it is imported
  without a date and reported as a warning).
* A *spec* is a grade's Standard LSL/USL per parameter plus its AIM LSL/USL (Standard.xlsx has one sheet for each;
  both can also be edited by hand).
* Charts are drawn per *spec* (e.g. "NI-Brass (Ni - 05) (5rs.)"), because the 5 Rs and
  10/20 Rs Ni-Brass coils are held to different limits.
"""
import csv
import difflib
import io
import json
import math
import re
import statistics
from collections import Counter, defaultdict
from functools import lru_cache

ELEMENTS = ["cu", "ni", "zn", "al", "mn", "fe", "pb", "sn", "si", "p", "s", "c"]
PARAMS = ELEMENTS + ["impurities", "total"]
PARAM_LABEL = {
    "cu": "Cu%", "ni": "Ni%", "zn": "Zn%", "al": "Al%", "mn": "Mn%", "fe": "Fe%", "pb": "Pb%",
    "sn": "Sn%", "si": "Si%", "p": "P%", "s": "S%", "c": "C%",
    "impurities": "Total Impurities%", "total": "Total%",
}
MAX_ROWS = 20000

_DATE_HEADERS = {"date", "castdate", "castingdate", "heatdate", "dateofcast", "dateofcasting", "dt", "castdt", "heatdt", "datetime", "timestamp"}


def _is_date_header(h):
    """h = normalised header (see _hnorm). Any header that says 'date' (Cast Date, Date of Analysis, Analysis Date, Sample Date,
    Reporting Date, Date/Time ...) is the cast date; 'update(d)' / 'validate' style words are not."""
    if h in _DATE_HEADERS:
        return True
    return "date" in h and not re.search(r"(update|validat|candidate|mandate)", h)
_ID_HEADERS = {"heatno", "heat", "coilno", "coil", "castno", "heatnumber", "heatnum"}
_ALLOY_HEADERS = {"alloy", "alloycode"}
_DENOM_HEADERS = {"denomination", "denom"}
_ANALYST_HEADERS = {"name", "analyst", "chemist", "analysedby", "analyzedby", "sampledby"}
_IMPURITY_HEADERS = {"impurities", "totalimpurities", "impurity"}
_TOTAL_HEADERS = {"total", "totalpct"}

# validation thresholds
TOTAL_WARN_DIFF = 0.02      # |Total% - sum(elements)| above this -> warning
TOTAL_ERR_DIFF = 1.0        # above this -> error (almost certainly a shifted decimal)

# Lab results are reported to 3 decimals (e.g. Pb 0.003, Cu 75.212): never round below that.
CHEM_DIGITS = 3

# Cpk rating for copper-base alloys (industry practice): >=1.67 excellent, >=1.33 capable, >=1.00 marginal.
CPK_EXCELLENT = 1.67
CPK_CAPABLE = 1.33
CPK_MARGINAL = 1.00
MAIN_MIN_MEAN = 1.0         # fallback when a spec names no alloying element: any element averaging >= 1 %


def cpk_rating(v):
    """'excellent' | 'capable' | 'marginal' | 'poor' | None (no value)."""
    if v is None:
        return None
    if v >= CPK_EXCELLENT:
        return "excellent"
    if v >= CPK_CAPABLE:
        return "capable"
    if v >= CPK_MARGINAL:
        return "marginal"
    return "poor"


def main_elements(limits, means=None):
    """Keys of the MAIN elements of a grade: copper (the base) plus every alloying element the spec
    requires a real minimum for (LSL > 0: Ni in Cu-Ni, Zn in brass, Al in Al-bronze ...). Elements with only
    a maximum, or a 0 minimum, are impurities. With no such element in the spec (or no spec) any element whose
    mean is >= 1 % counts. Order: Cu first, then by size of the mean, largest first."""
    limits = limits or {}
    means = means or {}
    main = {"cu"}
    for p in ELEMENTS:
        lsl, _usl = eff_limits(limits, p)
        if lsl is not None:
            main.add(p)
    if main == {"cu"}:
        for p in ELEMENTS:
            m = means.get(p)
            if m is not None and m >= MAIN_MIN_MEAN:
                main.add(p)
    return sorted(main, key=lambda p: (p != "cu", -(means.get(p) or 0), ELEMENTS.index(p)))


# ----------------------------------------------------------------------------- helpers
def _hnorm(h):
    return re.sub(r"[\s%._\-/()]+", "", str(h if h is not None else "").lower())


def norm_key(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def norm_heat(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", "", str(v)).upper()


def norm_denom(v):
    return re.sub(r"\s+", "", str(v or "")).upper()


def parse_number(v):
    """float | None (blank). Raises ValueError for text that is not a number."""
    if v is None or isinstance(v, bool):
        if isinstance(v, bool):
            raise ValueError("not a number")
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        if math.isnan(f) or math.isinf(f):
            raise ValueError("not a finite number")
        return f
    s = str(v).strip().replace("\u00a0", "")
    if s.endswith("%"):                      # "75.2 %" typed as text: the % sign is just a unit
        s = s[:-1].strip()
    if s in ("", "-", "--", "—", "NA", "N/A", "na", "n/a", "ND", "nd"):
        return None
    if "," in s and "." not in s and s.count(",") == 1:
        s = s.replace(",", ".")
    s = s.replace(",", "")
    try:
        f = float(s)
    except ValueError:
        raise ValueError("not a number")
    if math.isnan(f) or math.isinf(f):
        raise ValueError("not a finite number")
    return f


def heat_prefix(h):
    m = re.match(r"^[A-Z]+", h or "")
    return m.group(0) if m else ""


def heat_seq(h):
    m = re.search(r"(\d+)$", h or "")
    return int(m.group(1)) if m else -1


def order_key(rec):
    """Production order = heat number only: letters prefix, then the NUMERIC sequence (NBS999 < NBS1000)."""
    return (heat_prefix(rec.get("heat_no")), heat_seq(rec.get("heat_no")), rec.get("heat_no") or "")


# ----------------------------------------------------------------------------- cast date -> Month / Week / Quarter / FY
import datetime as _dt

_DATE_FMTS = ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%y", "%d/%m/%y", "%d-%m-%y", "%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d-%b-%Y", "%d-%b-%y",
              "%d %b %Y", "%d %B %Y", "%d-%B-%Y", "%d %b %y", "%d/%b/%Y", "%d/%b/%y", "%d.%b.%Y", "%d.%b.%y", "%d %b, %Y", "%b %d, %Y", "%B %d, %Y")
_TIME_TAIL = re.compile(r"(?:[T\s]+\d{1,2}[:.]\d{2}(?:[:.]\d{2}(?:\.\d+)?)?\s*(?:[AaPp][Mm])?\s*(?:Z|[+-]\d{2}:?\d{2})?)$")
_DATE_MIN_YEAR, _DATE_MAX_YEAR = 2000, 2100


def parse_cast_date(v):
    """Cast date cell -> (iso 'YYYY-MM-DD' | '', note). Plant files are day-first (02.04.2026, 04.04.26).
    note is '' when the value was read as typed, 'guessed' when a mistyped number was reconstructed (2505.2025 ->
    25.05.2025), 'unreadable' when nothing sensible could be read, 'blank' for an empty cell."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return "", "blank"
    if isinstance(v, bool):
        return "", "unreadable"
    d = None
    note = ""
    if isinstance(v, _dt.datetime):
        d = v.date()
    elif isinstance(v, _dt.date):
        d = v
    elif isinstance(v, (int, float)):
        f = float(v)
        if 30000 <= f <= 80000:                      # Excel serial day number
            d = (_dt.datetime(1899, 12, 30) + _dt.timedelta(days=int(f))).date()
        else:                                        # 25.05.2025 typed into a numeric cell and stored as 2505.2025
            m = re.match(r"^(\d{2})(\d{2})\.(\d{4})$", ("%.4f" % f))
            if m:
                try:
                    d = _dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
                    note = "guessed"
                except ValueError:
                    d = None
    else:
        t = re.sub(r"\s+", " ", str(v).strip())
        t = _TIME_TAIL.sub("", t).strip()            # '02.04.2026 10:30:00' / '2026-04-02T10:30:00' -> the date part
        for fmt in _DATE_FMTS:
            try:
                d = _dt.datetime.strptime(t, fmt).date()
                break
            except ValueError:
                continue
        if d is None:
            m = re.match(r"^(\d{2})(\d{2})\.(\d{4})$", t)          # '2505.2025' typed as text: 25.05.2025 (same guess as for a numeric cell)
            if m:
                try:
                    d = _dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
                    note = "guessed"
                except ValueError:
                    d = None
        if d is None:
            m = re.match(r"^(\d{4})(\d{2})(\d{2})$", t)          # 20260402
            if m:
                try:
                    d = _dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
                except ValueError:
                    d = None
    if d is None or not (_DATE_MIN_YEAR <= d.year <= _DATE_MAX_YEAR):
        return "", "unreadable"
    return d.isoformat(), note


@lru_cache(maxsize=12000)
def period_fields(iso):
    """ISO date -> (month, week, quarter, fy) with EXACTLY the labels the main dashboard uses
    (Apr-2026, Wk of 30-Mar-26 (Monday start), Q1 = Apr-Jun ..., FY 2026-27)."""
    try:
        d = _dt.date.fromisoformat(str(iso)[:10])
    except (TypeError, ValueError):
        return "", "", "", ""
    monday = d - _dt.timedelta(days=d.weekday())
    q = ((d.month - 4) % 12) // 3 + 1
    fy0 = d.year if d.month >= 4 else d.year - 1
    return d.strftime("%b-%Y"), "Wk of " + monday.strftime("%d-%b-%y"), f"Q{q}", f"FY {fy0}-{(fy0 + 1) % 100:02d}"


PERIOD_KEYS = ("month", "week", "quarter", "fy")


def filter_by_period(recs, flt):
    """flt: {month, week, quarter, fy} ('' / 'All' = no filter). Heats without a readable date drop out
    as soon as any period filter is on."""
    flt = {k: str((flt or {}).get(k) or "") for k in PERIOD_KEYS}
    flt = {k: ("" if v.strip().lower() == "all" else v.strip()) for k, v in flt.items()}
    if not any(flt.values()):
        return list(recs)
    wanted = tuple(flt[k] for k in PERIOD_KEYS)
    out = []
    for r in recs:
        pf = period_fields(r.get("cast_date"))
        if ((not wanted[0] or pf[0] == wanted[0]) and
            (not wanted[1] or pf[1] == wanted[1]) and
            (not wanted[2] or pf[2] == wanted[2]) and
            (not wanted[3] or pf[3] == wanted[3])):
            out.append(r)
    return out


def period_options(recs):
    """Distinct periods present in recs, newest first (weeks/months chronologically descending)."""
    mo, wk, qt, fy = {}, {}, set(), set()
    undated = 0
    for r in recs:
        iso = r.get("cast_date")
        m, w, q, f = period_fields(iso)
        if not m:
            undated += 1
            continue
        mo[m] = str(iso)[:7]
        # Sort key must be the ISO Monday date (chronological). Sorting the display label
        # "DD-Mon-YY" as text puts 30-Mar before 29-Jun before 27-Apr (wrong order).
        _d = _dt.date.fromisoformat(str(iso)[:10])
        wk[w] = (_d - _dt.timedelta(days=_d.weekday())).isoformat()
        qt.add(q)
        fy.add(f)
    return {"months": [k for k, _ in sorted(mo.items(), key=lambda kv: kv[1], reverse=True)],
            "weeks": [k for k, _ in sorted(wk.items(), key=lambda kv: kv[1], reverse=True)],
            "quarters": sorted(qt), "fys": sorted(fy, reverse=True),
            "undated": undated}


def period_options_cascade(recs, flt):
    """Options for each of the four period dropdowns, scoped to the OTHER three selections (same idea as the main dashboard's
    cascading filters): picking Month = Jun-2026 leaves only the weeks / quarter / FY that really occur in June, so a combination
    that would show 'no heats' can no longer be picked. A dropdown's own selection never narrows its own list."""
    flt = {k: str((flt or {}).get(k) or "").strip() for k in PERIOD_KEYS}
    flt = {k: ("" if v.lower() == "all" else v) for k, v in flt.items()}
    lists = {"month": "months", "week": "weeks", "quarter": "quarters", "fy": "fys"}
    out = {}
    for k, name in lists.items():
        others = {kk: vv for kk, vv in flt.items() if kk != k}
        out[name] = period_options(filter_by_period(recs, others))[name]
    return out


def previous_period_filter(recs, flt):
    """Return the immediately previous comparable single-period filter.

    Priority mirrors the main dashboard: Week > Month > Quarter > FY. The other
    period selectors are cleared so a selected Month/Week does not accidentally
    turn the comparison into an empty cross-period intersection. Returns None
    when no single comparable period is selected.
    """
    flt = {k: str((flt or {}).get(k) or "").strip() for k in PERIOD_KEYS}
    flt = {k: ("" if v.lower() == "all" else v) for k, v in flt.items()}

    if flt["week"]:
        m = re.match(r"^Wk of (\d{2}-[A-Za-z]{3}-\d{2})$", flt["week"])
        if not m:
            return None
        try:
            monday = _dt.datetime.strptime(m.group(1), "%d-%b-%y").date()
        except ValueError:
            return None
        prev = monday - _dt.timedelta(days=7)
        return {"month": "", "week": "Wk of " + prev.strftime("%d-%b-%y"), "quarter": "", "fy": ""}

    if flt["month"]:
        try:
            d = _dt.datetime.strptime(flt["month"], "%b-%Y").date().replace(day=1)
        except ValueError:
            return None
        prev = d - _dt.timedelta(days=1)
        return {"month": prev.strftime("%b-%Y"), "week": "", "quarter": "", "fy": ""}

    if flt["quarter"]:
        q = re.match(r"^Q([1-4])$", flt["quarter"], re.I)
        if not q:
            return None
        qn = int(q.group(1))
        fy = flt["fy"]
        if not fy:
            dated = []
            for r in recs:
                pf = period_fields(r.get("cast_date"))
                if pf[2] == flt["quarter"] and pf[3]:
                    dated.append(r.get("cast_date"))
            dated = [d for d in dated if d]
            if not dated:
                return None
            fy = period_fields(max(dated))[3]
        fm = re.match(r"^FY (\d{4})-(\d{2})$", fy, re.I)
        if not fm:
            return None
        fy0 = int(fm.group(1))
        if qn > 1:
            pq, pfy0 = qn - 1, fy0
        else:
            pq, pfy0 = 4, fy0 - 1
        return {"month": "", "week": "", "quarter": f"Q{pq}", "fy": f"FY {pfy0}-{(pfy0 + 1) % 100:02d}"}

    if flt["fy"]:
        fm = re.match(r"^FY (\d{4})-(\d{2})$", flt["fy"], re.I)
        if not fm:
            return None
        fy0 = int(fm.group(1)) - 1
        return {"month": "", "week": "", "quarter": "", "fy": f"FY {fy0}-{(fy0 + 1) % 100:02d}"}

    return None


# ----------------------------------------------------------------------------- reading files
def _map_header(cells):
    """header row -> {index: field}. Fields: heat_no, cast_date, alloy, denomination, analyst,
    <param>, hardness, conductivity, hf_no."""
    mapping = {}
    used = set()
    for i, raw in enumerate(cells):
        h = _hnorm(raw)
        if not h:
            continue
        field = None
        if _is_date_header(h):
            field = "cast_date"
        elif h in _ID_HEADERS:
            field = "heat_no"
        elif h in _ALLOY_HEADERS:
            field = "alloy"
        elif h in _DENOM_HEADERS:
            field = "denomination"
        elif h in _ANALYST_HEADERS:
            field = "analyst"
        elif h in _IMPURITY_HEADERS:
            field = "impurities"
        elif h in _TOTAL_HEADERS:
            field = "total"
        elif h in ELEMENTS:
            field = h
        elif h.startswith("hardness"):
            field = "hardness"
        elif h.startswith("conductivity"):
            field = "conductivity"
        elif h.startswith("hf"):
            field = "hf_no"
        if field and field not in used:
            mapping[i] = field
            used.add(field)
    return mapping


def _rows_from_table(sheet, table):
    """table: iterable of row tuples (first row = header). Returns (rows, note)."""
    it = iter(table)
    try:
        header = next(it)
    except StopIteration:
        return [], "sheet is empty"
    mapping = _map_header(header)
    if "heat_no" not in mapping.values():
        return [], "no heat / coil number column found"
    if not any(f in ELEMENTS for f in mapping.values()):
        return [], "no element (Cu%, Ni%, ...) columns found"
    rows = []
    for offset, cells in enumerate(it, start=2):
        if cells is None or all(c is None or str(c).strip() == "" for c in cells):
            continue
        raw = {}
        for i, field in mapping.items():
            raw[field] = cells[i] if i < len(cells) else None
        raw["_sheet"] = sheet
        raw["_row"] = offset
        rows.append(raw)
    return rows, ""


def read_chem_file(filename, data):
    """Parse .xlsx/.xlsm/.csv/.tsv bytes -> (raw_rows, sheet_notes).
    raw_rows keep original cell values; validate_rows() interprets them."""
    name = (filename or "").lower()
    rows, notes = [], []
    if name.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        try:
            for ws in wb.worksheets:
                got, note = _rows_from_table(ws.title, ws.iter_rows(values_only=True))
                if note:
                    notes.append({"sheet": ws.title, "rows": 0, "note": note})
                else:
                    notes.append({"sheet": ws.title, "rows": len(got), "note": ""})
                rows.extend(got)
                if len(rows) > MAX_ROWS:
                    raise ValueError(f"Import limited to {MAX_ROWS:,} rows per upload")
        finally:
            wb.close()
    elif name.endswith((".csv", ".tsv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        delim = "\t" if (name.endswith(".tsv") or text.count("\t") > text.count(",")) else ","
        table = list(csv.reader(io.StringIO(text), delimiter=delim))
        got, note = _rows_from_table("", table)
        notes.append({"sheet": "(file)", "rows": len(got), "note": note})
        rows.extend(got)
    else:
        raise ValueError("Unsupported file type. Upload .xlsx, .xlsm, .csv or .tsv")
    if not rows:
        detail = "; ".join(f"{n['sheet']}: {n['note']}" for n in notes if n["note"]) or "no data rows"
        raise ValueError("No chemistry rows found (" + detail + ")")
    if len(rows) > MAX_ROWS:
        raise ValueError(f"Import limited to {MAX_ROWS:,} rows per upload")
    return rows, notes


# ----------------------------------------------------------------------------- specs
def _limit_pair(lsl, usl):
    return [lsl, usl]


_AIM_SHEET_WORDS = ("aim", "target")


def _g(v):
    return "–" if v is None else ("%.6g" % v)


def _is_aim_sheet(title):
    t = norm_key(title)
    return any(w in t for w in _AIM_SHEET_WORDS)


def _parse_spec_table(table, label=""):
    """One Standard/AIM sheet -> (specs, issues). Columns: Alloy, Grade Descriptions, '<Param>% (LSL)', '<Param>% (USL)'."""
    if not table:
        raise ValueError(f"The {label or 'spec'} sheet is empty")
    header = table[0]
    col_alloy = col_desc = None
    limit_cols = {}          # (param, 'lsl'|'usl') -> index
    issues = []
    tag = f"[{label}] " if label else ""
    for i, raw in enumerate(header):
        h = str(raw or "").strip()
        if not h:
            continue
        hn = _hnorm(h)
        if hn in ("alloy", "alloycode"):
            col_alloy = i
        elif hn in ("gradedescriptions", "gradedescription", "grade", "description", "spec", "specname"):
            col_desc = i
        else:
            m = re.match(r"^(.*?)\s*\(\s*(LSL|USL)\s*\)\s*$", h, re.I)
            if m:
                p = _hnorm(m.group(1))
                if p in ("totalimpurities", "impurities", "impurity"):
                    p = "impurities"
                elif p in ("total", "totalpct"):
                    p = "total"
                if p in PARAMS:
                    limit_cols[(p, m.group(2).lower())] = i
                else:
                    issues.append({"row": 1, "severity": "warn", "code": "unknown_column",
                                   "message": f"{tag}Column '{h}' is not a known parameter and was ignored"})
    if col_desc is None:
        raise ValueError(f"No 'Grade Descriptions' column found{' on sheet ' + label if label else ''}")
    if not limit_cols:
        raise ValueError(f"No '(LSL)' / '(USL)' columns found{' on sheet ' + label if label else ''}")
    specs, seen = [], {}
    for rn, cells in enumerate(table[1:], start=2):
        if not cells or all(c is None or str(c).strip() == "" for c in cells):
            continue
        desc = str(cells[col_desc] if col_desc < len(cells) and cells[col_desc] is not None else "").strip()
        alloy = str(cells[col_alloy] if col_alloy is not None and col_alloy < len(cells) and cells[col_alloy] is not None else "").strip()
        if not desc:
            issues.append({"row": rn, "severity": "error", "code": "missing_description", "message": f"{tag}Grade description is blank"})
            continue
        limits, bad = {}, False
        for p in PARAMS:
            vals = {}
            for side in ("lsl", "usl"):
                idx = limit_cols.get((p, side))
                v = cells[idx] if idx is not None and idx < len(cells) else None
                try:
                    vals[side] = parse_number(v)
                except ValueError:
                    issues.append({"row": rn, "severity": "error", "code": "bad_limit",
                                   "message": f"{tag}{desc}: {PARAM_LABEL[p]} {side.upper()} is not a number"})
                    bad = True
            if bad:
                continue
            if vals["lsl"] is None and vals["usl"] is None:
                continue
            if vals["lsl"] is not None and vals["usl"] is not None and vals["lsl"] > vals["usl"]:
                issues.append({"row": rn, "severity": "error", "code": "lsl_gt_usl",
                               "message": f"{tag}{desc}: {PARAM_LABEL[p]} LSL ({vals['lsl']}) is greater than USL ({vals['usl']})"})
                bad = True
                continue
            limits[p] = _limit_pair(vals["lsl"], vals["usl"])
        if bad:
            continue
        if not limits:
            issues.append({"row": rn, "severity": "warn", "code": "no_limits", "message": f"{tag}{desc}: no limits filled in; skipped"})
            continue
        key = norm_key(desc)
        if key in seen:
            issues.append({"row": rn, "severity": "error", "code": "dup_description",
                           "message": f"{tag}'{desc}' is already defined at row {seen[key]}; this row was skipped"})
            continue
        seen[key] = rn
        specs.append({"alloy": alloy, "description": desc, "limits": limits, "row": rn})
    return specs, issues


def _read_spec_tables(filename, data):
    """-> [(sheet_title, table)]; csv/tsv give one unnamed table."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        try:
            return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
        finally:
            wb.close()
    if name.endswith((".csv", ".tsv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        delim = "\t" if (name.endswith(".tsv") or text.count("\t") > text.count(",")) else ","
        return [("", list(csv.reader(io.StringIO(text), delimiter=delim)))]
    raise ValueError("Unsupported file type. Upload .xlsx, .xlsm, .csv or .tsv")


def parse_spec_file(filename, data):
    """Standard.xlsx-style workbook -> (specs, issues).

    The workbook has one sheet with the STANDARD limits (customer/specification LSL/USL) and one sheet whose name
    contains 'AIM' with the plant's tighter AIM limits, both with the same layout. Each returned spec is
    {alloy, description, limits (Standard), aim (AIM limits, {} when the file has none for it), row}.
    A workbook with a single sheet gives Standard limits only."""
    tables = _read_spec_tables(filename, data)
    if not tables or not any(t for _, t in tables):
        raise ValueError("The spec file is empty")
    std_tabs = [(n, t) for n, t in tables if t and not _is_aim_sheet(n)]
    aim_tabs = [(n, t) for n, t in tables if t and _is_aim_sheet(n)]
    if not std_tabs:
        raise ValueError("No Standard-limits sheet found (only an AIM sheet was in the file)")
    pref = [x for x in std_tabs if "standard" in norm_key(x[0]) or "std" in norm_key(x[0])]
    std_name, std_table = (pref or std_tabs)[0]
    multi = len(tables) > 1
    specs, issues = _parse_spec_table(std_table, std_name if multi else "")
    for s in specs:
        s["aim"] = {}
    if aim_tabs:
        aim_name, aim_table = aim_tabs[0]
        aim_specs, aim_issues = _parse_spec_table(aim_table, aim_name)
        issues.extend(aim_issues)
        by_key = {norm_key(s["description"]): s for s in specs}
        for a in aim_specs:
            akey = norm_key(a["description"])
            s = by_key.get(akey)
            if not s:
                # Grade names are often typed slightly differently on the two sheets ("Brass (5rs.)" vs "Brass (5rs.) AIM"):
                # accept a match when exactly ONE Standard grade contains / is contained in the AIM name and has no AIM yet.
                cand = [v for k, v in by_key.items() if len(min(k, akey, key=len)) >= 4 and (k in akey or akey in k) and not v["aim"]]
                if len(cand) == 1:
                    s = cand[0]
                    issues.append({"row": a["row"], "severity": "info", "code": "aim_loose_match",
                                   "message": f"[{aim_name}] '{a['description']}' was matched to the Standard grade '{s['description']}' by similar name"})
            if not s:
                issues.append({"row": a["row"], "severity": "warn", "code": "aim_without_standard",
                               "message": f"[{aim_name}] '{a['description']}' has no matching row on the Standard sheet; its AIM limits were ignored"})
                continue
            s["aim"] = a["limits"]
            for p, (alsl, ausl) in a["limits"].items():
                slsl, susl = eff_limits(s["limits"], p)
                if (alsl is not None and slsl is not None and alsl < slsl - 1e-9) or (ausl is not None and susl is not None and ausl > susl + 1e-9):
                    issues.append({"row": a["row"], "severity": "warn", "code": "aim_outside_standard",
                                   "message": f"[{aim_name}] {a['description']}: {PARAM_LABEL[p]} AIM {_g(alsl)}–{_g(ausl)} is wider than the Standard limits {_g(slsl)}–{_g(susl)}"})
        for s in specs:
            if not s["aim"]:
                issues.append({"row": s["row"], "severity": "info", "code": "no_aim",
                               "message": f"{s['description']}: no AIM row found; only Standard limits will be used"})
    if not specs:
        raise ValueError("No usable spec rows found")
    return specs, issues


def validate_aim_payload(aim):
    """Clean manually edited AIM limits. Returns (clean_aim, error_or_None); empty is allowed."""
    if aim in (None, {}, []):
        return {}, None
    if not isinstance(aim, dict):
        return None, "AIM limits are malformed"
    clean = {}
    for p, pair in aim.items():
        if p not in PARAMS:
            return None, f"Unknown parameter '{p}' in the AIM limits"
        if not isinstance(pair, (list, tuple)) or len(pair) != 2:
            return None, f"AIM {PARAM_LABEL[p]}: limits must be [LSL, USL]"
        try:
            lsl, usl = parse_number(pair[0]), parse_number(pair[1])
        except ValueError:
            return None, f"AIM {PARAM_LABEL[p]}: LSL/USL must be numbers"
        if lsl is None and usl is None:
            continue
        if lsl is not None and usl is not None and lsl > usl:
            return None, f"AIM {PARAM_LABEL[p]}: LSL is greater than USL"
        clean[p] = [lsl, usl]
    return clean, None


def validate_spec_payload(alloy, description, limits):
    """Clean a manually edited spec. Returns (clean_limits, error_or_None)."""
    description = str(description or "").strip()
    if not description:
        return None, "Grade description is required"
    if len(description) > 120 or len(str(alloy or "")) > 40:
        return None, "Description or alloy code is too long"
    if not isinstance(limits, dict) or not limits:
        return None, "Enter at least one LSL/USL"
    clean = {}
    for p, pair in limits.items():
        if p not in PARAMS:
            return None, f"Unknown parameter '{p}'"
        if not isinstance(pair, (list, tuple)) or len(pair) != 2:
            return None, f"{PARAM_LABEL[p]}: limits must be [LSL, USL]"
        try:
            lsl, usl = parse_number(pair[0]), parse_number(pair[1])
        except ValueError:
            return None, f"{PARAM_LABEL[p]}: LSL/USL must be numbers"
        if lsl is None and usl is None:
            continue
        if lsl is not None and usl is not None and lsl > usl:
            return None, f"{PARAM_LABEL[p]}: LSL is greater than USL"
        clean[p] = [lsl, usl]
    if not clean:
        return None, "Enter at least one LSL/USL"
    return clean, None


def sheet_alloy_conflict(spec, alloy):
    """True when a heat's Alloy column and the alloy code of the spec its sheet name points to are both
    filled in and different (so the sheet name must not decide the spec)."""
    sa, an = norm_key((spec or {}).get("alloy")), norm_key(alloy)
    return bool(sa and an and sa != an)


def sheet_named_spec(specs, sheet):
    """The spec whose description equals the sheet name exactly (None when there is none)."""
    sn = norm_key(sheet)
    if not sn:
        return None
    return next((s for s in specs or [] if norm_key(s["description"]) == sn), None)


def resolve_spec(specs, alloy, sheet="", denom=""):
    """Pick the spec that applies to a heat. specs: [{alloy, description, limits}, ...]"""
    if not specs:
        return None
    sn = norm_key(sheet)
    an = norm_key(alloy)
    if sn:
        for s in specs:
            if norm_key(s["description"]) == sn:
                # A sheet named exactly like a grade only decides when it does not contradict the heat's own
                # Alloy column. If both are present and differ, the sheet name is not trusted (a wrongly named
                # sheet would otherwise put the heat on the wrong limits and give a wrong Cpk): the Alloy
                # column decides instead, and when it cannot, the heat stays unassigned rather than guessed.
                if not sheet_alloy_conflict(s, alloy):
                    return s
                sn = ""
                break
    cands = [s for s in specs if an and norm_key(s.get("alloy")) == an]
    if len(cands) == 1:
        return cands[0]
    if len(cands) > 1:
        dn = norm_key(denom)
        if dn:
            by_d = [s for s in cands if dn in norm_key(s["description"])]
            if len(by_d) == 1:
                return by_d[0]
        if sn:
            by_s = [s for s in cands if norm_key(s["description"]) in sn or sn in norm_key(s["description"])]
            if len(by_s) == 1:
                return by_s[0]
        return None
    if sn:
        by_s = [s for s in specs if len(sn) >= 6 and (sn in norm_key(s["description"]) or norm_key(s["description"]) in sn)]
        if len(by_s) == 1:
            return by_s[0]
    return None


def eff_limits(limits, param):
    """(lsl, usl) actually enforced. A 0 LSL is the natural bound for impurities, so it is ignored."""
    pair = (limits or {}).get(param)
    if not pair:
        return None, None
    lsl, usl = pair
    if lsl is not None and lsl <= 0:
        lsl = None
    return lsl, usl


def spec_violations(rec, limits):
    """[{param, value, side, limit}] for every parameter outside its limits."""
    out = []
    for p, pair in (limits or {}).items():
        v = rec.get(p)
        if v is None:
            continue
        lsl, usl = eff_limits(limits, p)
        cmpv = round(v, CHEM_DIGITS) if p == "total" else v      # Total% is a computed sum; ignore 4th-decimal noise, keep all 3 reported decimals
        if lsl is not None and cmpv < lsl - 1e-9:
            out.append({"param": p, "value": v, "side": "below", "limit": lsl})
        elif usl is not None and cmpv > usl + 1e-9:
            out.append({"param": p, "value": v, "side": "above", "limit": usl})
    return out


# ----------------------------------------------------------------------------- import validation
def _sheet_denoms(sheet):
    nums = set()
    for m in re.finditer(r"(\d+(?:\s*[-&]\s*\d+)*)\s*rs\b", str(sheet or ""), re.I):
        nums.update(re.findall(r"\d+", m.group(1)))
    return nums


def _num_changed(a, b):
    if a is None and b is None:
        return False
    if a is None or b is None:
        return True
    return abs(a - b) > 1e-9


def _record_diff(old, new):
    changes = []
    for f in ("sheet", "cast_date", "alloy", "denomination", "analyst", "hardness", "hf_no"):
        if str(old.get(f) or "") != str(new.get(f) or ""):
            changes.append({"field": f, "old": old.get(f) or "", "new": new.get(f) or ""})
    for p in PARAMS:
        if _num_changed(old.get(p), new.get(p)):
            changes.append({"field": PARAM_LABEL[p], "old": old.get(p), "new": new.get(p)})
    if _num_changed(old.get("conductivity"), new.get("conductivity")):
        changes.append({"field": "Conductivity", "old": old.get("conductivity"), "new": new.get("conductivity")})
    return changes


_KEEP_IF_BLANK = ("sheet", "cast_date", "alloy", "denomination", "analyst", "hardness", "hf_no")


def _merge_blank_from_existing(rec, old):
    """A re-import must never ERASE identifying details: if the new file leaves an analyst/sheet/... blank
    (e.g. a CSV without sheet names or analyst) the stored value is kept.
    Chemistry numbers are different: the new file is the truth for them, so they are overwritten."""
    for f in _KEEP_IF_BLANK:
        if not rec.get(f) and old.get(f):
            rec[f] = old[f]
    if rec.get("conductivity") is None and old.get("conductivity") is not None:
        rec["conductivity"] = old["conductivity"]


def validate_rows(raw_rows, existing=None, specs=None, max_issues=400):
    """Interpret raw rows and run every import check.

    existing: {heat_no: record} already in the database.
    Returns dict: records (to write, each has _status 'new'|'update'), unchanged, duplicates, errors,
    warnings, issues, issue_counts, updated_details, oos_heats, unresolved_sheets.
    Errors exclude the row; warnings keep it (but are shown so the source can be fixed).
    """
    existing = existing or {}
    specs = specs or []
    issues = []
    counts = Counter()
    sev_counts = Counter()

    def issue(sev, code, msg, r=None, heat=""):
        counts[code] += 1
        sev_counts[sev] += 1
        if len(issues) < max_issues:
            issues.append({"sheet": (r or {}).get("_sheet", ""), "row": (r or {}).get("_row"), "heat_no": heat,
                           "severity": sev, "code": code, "message": msg})

    built = []
    for r in raw_rows:
        heat = norm_heat(r.get("heat_no"))
        rec = {"heat_no": heat, "sheet": r.get("_sheet", ""), "src_row": r.get("_row"), "_row": r.get("_row")}
        bad = False
        if not heat:
            issue("error", "missing_heat", "Heat / Coil No. is blank", r)
            continue
        if len(heat) > 40 or not re.match(r"^[A-Z0-9][A-Z0-9\-_/]*$", heat):
            issue("error", "bad_heat", f"Heat No. '{heat}' has unusual characters", r, heat)
            continue
        rec["alloy"] = str(r.get("alloy") or "").strip()
        rec["cast_date"], dnote = parse_cast_date(r.get("cast_date"))
        if dnote == "unreadable":
            issue("warn", "bad_date", f"Date '{r.get('cast_date')}' could not be read; the heat is imported without a date (it will not appear under Month/Week/Quarter/FY filters)", r, heat)
        elif dnote == "guessed":
            issue("warn", "date_guess", f"Date typed as the number {r.get('cast_date')}; read as {rec['cast_date']} (dd.mm.yyyy). Please confirm it in the source file", r, heat)
        elif dnote == "blank":
            issue("info", "no_date", "Date is blank; the heat is imported without a date", r, heat)
        if rec["cast_date"] and rec["cast_date"] > _dt.date.today().isoformat():
            issue("warn", "future_date", f"Date {rec['cast_date']} is in the future; check the day/month/year in the source (the heat is still imported with this date)", r, heat)
        rec["denomination"] = norm_denom(r.get("denomination"))
        rec["analyst"] = re.sub(r"\s+", " ", str(r.get("analyst") or "")).strip().upper()
        rec["hardness"] = "" if r.get("hardness") is None else str(r.get("hardness")).strip()
        rec["hf_no"] = "" if r.get("hf_no") is None else str(r.get("hf_no")).strip()
        try:
            rec["conductivity"] = parse_number(r.get("conductivity"))
        except ValueError:
            rec["conductivity"] = None
        # numbers
        for p in PARAMS:
            try:
                v = parse_number(r.get(p))
            except ValueError:
                issue("error", "bad_number", f"{PARAM_LABEL[p]} '{r.get(p)}' is not a number", r, heat)
                bad = True
                break
            if v is not None:
                lo, hi = (90.0, 110.0) if p == "total" else (0.0, 100.0)
                if v < lo or v > hi:
                    issue("error", "out_of_range", f"{PARAM_LABEL[p]} = {v:g} is outside the possible {lo:g}–{hi:g}% range", r, heat)
                    bad = True
                    break
            rec[p] = v
        if bad:
            continue
        el_vals = [rec[p] for p in ELEMENTS if rec.get(p) is not None]
        if not el_vals:
            issue("error", "no_values", "No chemistry values in this row", r, heat)
            continue
        s = sum(el_vals)
        if rec.get("total") is None:
            rec["total"] = s
        else:
            diff = abs(rec["total"] - s)
            if diff > TOTAL_ERR_DIFF:
                issue("error", "total_mismatch", f"Total% {rec['total']:.4f} is far from the sum of elements {s:.4f} (likely a typing error)", r, heat)
                continue
            if diff > TOTAL_WARN_DIFF:
                issue("warn", "total_mismatch", f"Total% {rec['total']:.4f} differs from the sum of elements {s:.4f}", r, heat)
        built.append(rec)

    # ---- duplicate heat_no inside the file
    by_heat = defaultdict(list)
    for rec in built:
        by_heat[rec["heat_no"]].append(rec)
    keep, dup_skipped = [], 0
    for heat, group in by_heat.items():
        if len(group) == 1:
            keep.append(group[0])
            continue
        first = group[0]
        if all(not _record_diff(first, g) for g in group[1:]):
            dup_skipped += len(group) - 1
            for g in group[1:]:
                issue("info", "dup_identical", f"Heat {heat} repeats row {first['_row']} with identical values; skipped", {"_sheet": g["sheet"], "_row": g["_row"]}, heat)
            keep.append(first)
        else:
            where = ", ".join(f"{g['sheet'] or 'file'} row {g['_row']}" for g in group)
            for g in group:
                issue("error", "dup_conflict", f"Heat {heat} appears {len(group)} times with different values ({where}); none imported, fix the source", {"_sheet": g["sheet"], "_row": g["_row"]}, heat)
            dup_skipped += len(group)
    built = keep

    # ---- a re-import must never erase stored identifying details; do this BEFORE spec matching so a file
    # without a sheet/alloy column still resolves to the same spec as the stored heat
    for rec in built:
        old_rec = existing.get(rec["heat_no"])
        if old_rec is not None:
            _merge_blank_from_existing(rec, old_rec)

    # ---- learned cross-checks (alloy per heat prefix, analyst spelling)
    pref = defaultdict(Counter)
    for rec in list(built) + list(existing.values()):
        if rec.get("alloy"):
            pref[heat_prefix(rec["heat_no"])][rec["alloy"]] += 1
    analysts = Counter()
    in_file = {r["heat_no"] for r in built}
    for rec in built:
        if rec.get("analyst"):
            analysts[rec["analyst"]] += 1
    for hn, rec in existing.items():          # heats already stored (a re-imported heat is counted once, from the file)
        if hn not in in_file and rec.get("analyst"):
            analysts[rec["analyst"]] += 1
    typo_map = {}
    for nm, cnt in analysts.items():
        if cnt > 3:
            continue
        best = None
        for other, ocnt in analysts.items():
            if other != nm and ocnt >= max(5, cnt * 5):
                ratio = difflib.SequenceMatcher(None, nm, other).ratio()
                if ratio >= 0.8 and (best is None or ratio > best[1]):
                    best = (other, ratio)
        if best:
            typo_map[nm] = best[0]

    unresolved = defaultdict(int)
    typo_seen = set()
    for rec in built:
        r = {"_sheet": rec["sheet"], "_row": rec["_row"]}
        heat = rec["heat_no"]
        pc = pref.get(heat_prefix(heat))
        if pc and rec.get("alloy"):
            top, n_top = pc.most_common(1)[0]
            total_n = sum(pc.values())
            if rec["alloy"] != top and total_n >= 5 and n_top / total_n >= 0.8:
                issue("warn", "alloy_prefix", f"Alloy '{rec['alloy']}' does not match the usual '{top}' for heats starting {heat_prefix(heat)}", r, heat)
        sd = _sheet_denoms(rec["sheet"])
        dn = re.match(r"^(\d+)RS$", rec.get("denomination") or "")
        # Excel truncates sheet names at 31 characters, so a name that long may have lost its tail
        # ("... (10rs. & 20r"); the denominations read from it are then incomplete and must not be trusted.
        if sd and dn and dn.group(1) not in sd and len(str(rec["sheet"] or "")) < 30:
            issue("warn", "denom_sheet", f"Denomination {rec['denomination']} is on a sheet for {'/'.join(sorted(sd, key=int))} Rs", r, heat)
        if rec["analyst"] in typo_map and rec["analyst"] not in typo_seen:
            typo_seen.add(rec["analyst"])
            issue("warn", "analyst_typo", f"Name '{rec['analyst']}' ({analysts[rec['analyst']]}×) looks like a typo of '{typo_map[rec['analyst']]}'", r, heat)
        spec = resolve_spec(specs, rec.get("alloy"), rec.get("sheet"), rec.get("denomination")) if specs else None
        rec["_spec"] = spec["description"] if spec else ""
        named = sheet_named_spec(specs, rec.get("sheet")) if specs else None
        if named and sheet_alloy_conflict(named, rec.get("alloy")):
            issue("warn", "sheet_alloy", f"Sheet '{rec['sheet']}' is named like grade '{named['description']}' (alloy {named.get('alloy')}) but this heat's Alloy column says '{rec.get('alloy')}'; "
                  + (f"the sheet name was ignored and the Alloy column chose '{spec['description']}'" if spec else "the sheet name was ignored and no spec could be chosen from the Alloy column, so the heat has no limits"), r, heat)
        if specs and not spec:
            unresolved[rec["sheet"] or rec.get("alloy") or "(unknown)"] += 1
        elif spec:
            for p, pair in spec["limits"].items():
                v = rec.get(p)
                lsl, usl = eff_limits(spec["limits"], p)
                if v is None or p == "total" or lsl is None or usl is None:
                    continue
                w = usl - lsl
                if v < lsl - w or v > usl + w:
                    issue("warn", "far_from_spec", f"{PARAM_LABEL[p]} = {v:g} is more than one full spec width outside {lsl:g}–{usl:g}; check it is not a typing error", r, heat)
    for k, n in unresolved.items():
        issue("warn", "no_spec", f"{n} heat(s) on '{k}' have no matching spec; charts will show them without LSL/USL until a spec is added", None, "")

    # ---- new / update / unchanged vs database
    records, unchanged, updated_details, oos = [], 0, [], 0
    for rec in built:
        old = existing.get(rec["heat_no"])
        if old is None:
            rec["_status"] = "new"
            records.append(rec)
        else:
            ch = _record_diff(old, rec)
            if ch:
                rec["_status"] = "update"
                records.append(rec)
                if len(updated_details) < 25:
                    updated_details.append({"heat_no": rec["heat_no"], "changes": ch[:8]})
            else:
                unchanged += 1
        spec = next((s for s in specs if s["description"] == rec.get("_spec")), None)
        if spec and spec_violations(rec, spec["limits"]):
            oos += 1
    return {
        "detected": len(raw_rows),
        "records": records,
        "new": sum(1 for r in records if r["_status"] == "new"),
        "updated": sum(1 for r in records if r["_status"] == "update"),
        "unchanged": unchanged,
        "duplicates": dup_skipped,
        "errors": sev_counts["error"],
        "warnings": sev_counts["warn"],
        "issues": issues,
        "issue_counts": dict(counts),
        "updated_details": updated_details,
        "oos_heats": oos,
        "unresolved_sheets": dict(unresolved),
    }


# ----------------------------------------------------------------------------- SPC maths
D2 = 1.128          # d2 for n=2 (moving range of 2)
E2 = 3 / D2         # 2.6596 ~ 2.66
D4 = 3.267          # MR-chart UCL factor


def _stdev(values):
    return statistics.stdev(values) if len(values) > 1 else 0.0


def imr(values):
    """Individuals / moving-range chart parameters (sigma from MRbar/d2)."""
    n = len(values)
    if n < 2:
        return None
    mr = [abs(values[i] - values[i - 1]) for i in range(1, n)]
    mrbar = sum(mr) / len(mr)
    cl = sum(values) / n
    sigma = mrbar / D2
    return {
        "n": n, "cl": cl, "mrbar": mrbar, "sigma_within": sigma,
        "ucl": cl + 3 * sigma, "lcl": cl - 3 * sigma,
        "mr": mr, "mr_ucl": D4 * mrbar, "mr_lcl": 0.0,
        "flat": mrbar == 0,
    }


def capability(values, lsl, usl, sigma_within, count_digits=None):
    """Cp/Cpk (within sigma) and Pp/Ppk (overall sigma). One-sided specs give Cpk/Ppk only."""
    n = len(values)
    res = {"n": n, "mean": None, "sigma_overall": None, "sigma_within": sigma_within, "cp": None, "cpk": None,
           "pp": None, "ppk": None, "lsl": lsl, "usl": usl, "n_below": 0, "n_above": 0, "ppm_observed": None,
           "note": ""}
    if n == 0:
        return res
    mean = sum(values) / n
    so = _stdev(values)
    res["mean"], res["sigma_overall"] = mean, so
    cv = [round(v, count_digits) for v in values] if count_digits is not None else values
    if lsl is not None:
        res["n_below"] = sum(1 for v in cv if v < lsl - 1e-9)
    if usl is not None:
        res["n_above"] = sum(1 for v in cv if v > usl + 1e-9)
    res["ppm_observed"] = (res["n_below"] + res["n_above"]) / n * 1e6
    if lsl is None and usl is None:
        res["note"] = "No limits set for this parameter"
        return res
    if n < 2 or not so or not sigma_within:
        res["note"] = "No variation in the data (all values identical), so capability cannot be computed" if n >= 2 else "Need at least 2 heats"
        return res

    def idx(s):
        parts = []
        if usl is not None:
            parts.append((usl - mean) / (3 * s))
        if lsl is not None:
            parts.append((mean - lsl) / (3 * s))
        return min(parts)
    res["cpk"], res["ppk"] = idx(sigma_within), idx(so)
    if lsl is not None and usl is not None:
        res["cp"] = (usl - lsl) / (6 * sigma_within)
        res["pp"] = (usl - lsl) / (6 * so)
    else:
        res["note"] = "One-sided specification: Cp/Pp are not defined, only Cpk/Ppk"
    if n < 30:
        res["note"] = (res["note"] + "; " if res["note"] else "") + f"Only {n} heats: treat capability as indicative"
    return res


def histogram(values, lsl, usl, nbins=None):
    """Bins cover the data range; axis range is widened to show LSL/USL when they are not far off-scale."""
    n = len(values)
    if n == 0:
        return {"bins": [], "xmin": 0, "xmax": 1, "lsl_off": False, "usl_off": False, "width": 0}
    dmin, dmax = min(values), max(values)
    span = dmax - dmin
    k = nbins or min(40, max(8, int(math.sqrt(n))))
    if span == 0:
        pad = abs(dmin) * 0.01 or 0.01
        bins = [{"x0": dmin - pad, "x1": dmin + pad, "n": n}]
        width = 2 * pad
        lo, hi = dmin - pad, dmin + pad
    else:
        width = span / k
        counts = [0] * k
        for v in values:
            i = min(k - 1, int((v - dmin) / width))
            counts[i] += 1
        bins = [{"x0": dmin + i * width, "x1": dmin + (i + 1) * width, "n": counts[i]} for i in range(k)]
        lo, hi = dmin, dmax
    ref = span or width
    lsl_off = usl_off = False
    xmin, xmax = lo, hi
    if lsl is not None:
        if lo - lsl <= 4 * ref:
            xmin = min(xmin, lsl)
        else:
            lsl_off = True
    if usl is not None:
        if usl - hi <= 4 * ref:
            xmax = max(xmax, usl)
        else:
            usl_off = True
    pad = (xmax - xmin) * 0.04 or 0.01
    return {"bins": bins, "xmin": xmin - pad, "xmax": xmax + pad, "lsl_off": lsl_off, "usl_off": usl_off, "width": width}


def analyse_param(points, limits, param):
    """points: [{value, ...}] in production order (already filtered to non-null). Returns chart payload pieces."""
    values = [p["value"] for p in points]
    lsl, usl = eff_limits(limits, param)
    ch = imr(values)
    payload = {"n": len(values), "lsl": lsl, "usl": usl, "imr": None,
               "mr_ooc": [], "capability": capability(values, lsl, usl, ch["sigma_within"] if ch else None, CHEM_DIGITS if param == "total" else None),
               "histogram": histogram(values, lsl, usl), "warnings": []}
    if ch is None:
        payload["warnings"].append("Need at least 2 heats to draw control charts")
        return payload
    payload["imr"] = {k: v for k, v in ch.items() if k != "mr"}
    payload["mr"] = ch["mr"]
    if ch["flat"]:
        payload["warnings"].append("All values are identical, so there are no control limits to draw")
    else:
        payload["mr_ooc"] = [i + 1 for i, m in enumerate(ch["mr"]) if m > ch["mr_ucl"]]
    if param == "total":
        cap = payload["capability"]
        cap["cp"] = cap["cpk"] = cap["pp"] = cap["ppk"] = None
        cap["note"] = "Total% is the sum of the reported elements, so Cp/Cpk are not meaningful for it (only the out-of-spec count is)"
    if len(values) < 20:
        payload["warnings"].append(f"Only {len(values)} heats: control limits from fewer than 20 points are provisional")
    return payload


def summarize_disposition(rows):
    """rows: disposition rows of ONE heat -> dict used by the SPC UI."""
    n = len(rows)
    qty = sum(float(r.get("output_weight") or 0) for r in rows)
    rej_rows = [r for r in rows if str(r.get("quality_decision") or "").strip().upper() == "REJECT"]
    rej_qty = sum(float(r.get("output_weight") or 0) for r in rej_rows)
    defects = Counter()
    decisions = Counter()
    for r in rows:
        md = str(r.get("main_defect") or "").strip().upper()
        if md and md != "NO DEFECT":
            defects[md] += 1
        decisions[str(r.get("quality_decision") or "").strip().upper() or "—"] += 1
    defect_coils = sum(defects.values())
    return {
        "coils": n, "qty_mt": round(qty, 3), "reject_coils": len(rej_rows), "reject_mt": round(rej_qty, 3),
        "reject_pct": round(rej_qty / qty * 100, 2) if qty > 0 else 0.0,
        "defect_coils": defect_coils, "defect_pct": round(defect_coils / n * 100, 1) if n else 0.0,
        "top_defects": [{"defect": d, "coils": c} for d, c in defects.most_common(3)],
        "decisions": dict(decisions),
    }


def build_spc_view(records, spec, param, disp_by_heat, last_n=0, max_oos=500, period=None):
    """Everything the Chemistry SPC tab needs for one spec group.

    records: chem records for this spec (dicts with heat_no, cast_date, analyst, params...); ordered here by heat number.
    disp_by_heat: {heat_no: summarize_disposition(...)} for heats that exist in disposition.
    period: optional {month, week, quarter, fy} filter on the cast date (applied before 'last N heats').
    """
    recs = sorted(filter_by_period(records, period), key=order_key)
    if last_n and last_n > 0:
        recs = recs[-last_n:]
    limits = (spec or {}).get("limits") or {}
    aim_lsl, aim_usl = eff_limits((spec or {}).get("aim") or {}, param)
    pts = [{"value": r[param]} for r in recs if r.get(param) is not None]
    a = analyse_param(pts, limits, param)
    # Compute spec violations once per selected heat. The old implementation
    # evaluated the full parameter-limit matrix twice and also injected a
    # temporary ``value`` key into the revision-cached source dictionaries.
    # Keeping the source snapshot immutable is safer under concurrent viewers.
    viol_by_heat = {r["heat_no"]: (spec_violations(r, limits) if limits else []) for r in recs}
    series = []
    for i, r in enumerate([r for r in recs if r.get(param) is not None]):
        viol = viol_by_heat.get(r["heat_no"], [])
        prm_viol = [v for v in viol if v["param"] == param]
        series.append({
            "i": i + 1, "heat_no": r["heat_no"], "cast_date": r.get("cast_date") or "", "value": r[param], "analyst": r.get("analyst", ""),
            "oos": bool(prm_viol), "heat_oos": bool(viol),
            "mr": a["mr"][i - 1] if i > 0 and a.get("mr") else None,
            "disp": disp_by_heat.get(r["heat_no"]),
        })
    # out-of-spec heats across ALL parameters that have limits
    oos_rows = []
    for r in recs:
        viol = viol_by_heat.get(r["heat_no"], [])
        if viol:
            oos_rows.append({"heat_no": r["heat_no"], "analyst": r.get("analyst", ""),
                             "violations": [{"param": PARAM_LABEL[v["param"]], "key": v["param"], "value": v["value"],
                                             "side": v["side"], "limit": v["limit"]} for v in viol],
                             "disp": disp_by_heat.get(r["heat_no"])})
    # in-spec vs out-of-spec comparison (only heats that have disposition data)
    oos_heats = {o["heat_no"] for o in oos_rows}

    def agg(names):
        ds = [disp_by_heat[h] for h in names if h in disp_by_heat]
        if not ds:
            return {"heats": 0, "coils": 0, "reject_pct": None, "defect_pct": None}
        qty = sum(d["qty_mt"] for d in ds)
        coils = sum(d["coils"] for d in ds)
        return {"heats": len(ds), "coils": coils,
                "reject_pct": round(sum(d["reject_mt"] for d in ds) / qty * 100, 2) if qty else None,
                "defect_pct": round(sum(d["defect_coils"] for d in ds) / coils * 100, 1) if coils else None}
    all_heats = [r["heat_no"] for r in recs]
    # With no limits nothing can be out of spec, but nothing has been shown to be IN spec either.
    in_spec = [h for h in all_heats if h not in oos_heats] if limits else []
    return {
        "param": param, "n": a["n"], "n_heats": len(recs),
        "lsl": a["lsl"], "usl": a["usl"], "aim_lsl": aim_lsl, "aim_usl": aim_usl,
        "imr": a["imr"], "mr": a.get("mr", []), "mr_ooc": a["mr_ooc"],
        "capability": a["capability"], "histogram": a["histogram"], "warnings": a["warnings"],
        "series": series,
        "oos_heats": oos_rows[:max_oos], "oos_total": len(oos_rows),
        "summary": {
            "heats": len(recs), "heats_with_disposition": sum(1 for h in all_heats if h in disp_by_heat),
            "oos_heats": len(oos_rows),
            "has_limits": bool(limits),
            "compare": {"in_spec": agg(in_spec), "out_of_spec": agg(sorted(oos_heats))},
        },
    }
