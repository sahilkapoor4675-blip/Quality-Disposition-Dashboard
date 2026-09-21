import sys, tempfile
from pathlib import Path
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parent
if not (ROOT / 'server.py').exists():   # works from the repo root or from a tests/ subfolder
    ROOT = ROOT.parent
sys.path.insert(0,str(ROOT))

from reports import _excel_report, _pdf_report, _pptx_report

N=25

def rows(prefix):
    return [{'name':f'{prefix}_{i}','output_qty':i+1,'coils':i+2,'defect_coils':1,'defect_pct':0.01,'reject_qty':0.1,'reject_pct_qty':0.001} for i in range(N)]

wc=rows('WC'); gr=rows('GR')
trend=[{'name':f'P{i}','coils':10,'output_qty':10,'defect_coils':1,'defect_pct':.1,'reject_qty':.1,'reject_pct_qty':.01,'first_pass_yield_pct':.9} for i in range(N)]
reg=[{'rank':i+1,'defect':f'D{i}','records':1,'qty':float(i+1),'pct_records':1/N,'cum_pct':min(1,(i+1)/N)} for i in range(N)]
causes={k:[f'Cause {i}' for i in range(8)] for k in ('man','machine','material','method','measurement','environment')}
rca={k:[{'root_cause':f'RC{i}','action':'Action','preventive_action':'Prevent','role':'Role','responsibility':'Owner','why_chain':['Why 1','Why 2']} for i in range(3)] for k in causes}
payload={
 'filters': {'month':'All','week':'All','quarter':'All','financial_year':'All','work_center':'All','grade':'All','main_defect':'All','defect_intensity':'All','quality_decision':'All'},
 'kpis': {'kpis':[{'label':f'K{i}','value':i,'fmt':''} for i in range(12)],'decision_table':[{'decision':f'DEC_{i}','qty':i+1} for i in range(6)],'intensity_table':[{'intensity':f'I{i}','qty':i+1,'coils':10} for i in range(8)]},
 'defects': {'register':reg,'pareto':reg[:10],'register_total':{'records':N,'qty':sum(range(1,N+1)),'pct_records':1.0}},
 'wcg': {'by_work_center':wc,'by_grade':gr,'total_work_center':{'coils':999,'output_qty':999,'defect_coils':1,'defect_pct':.01,'reject_qty':.1,'reject_pct_qty':.001},'total_grade':{'coils':999,'output_qty':999,'defect_coils':1,'defect_pct':.01,'reject_qty':.1,'reject_pct_qty':.001}},
 'monthly': {'rows':trend,'total':{}}, 'period': {'rows':trend,'total':{}}, 'quarterly': {'rows':trend[:10],'total':{}}, 'yearly': {'rows':trend[:5],'total':{}},
 'intel': {'early_warnings':[],'recurring_patterns':[],'kpi_ranking':[],'risk_matrix':{'work_centers':wc,'grades':gr},'health_score':{'reasons':[]}},
 'root_cause': {'defect':'D0','rows':[{'grade':'A','work_center':'WC','heat_no':'H','batch_no':'B','output_weight':1.0} for _ in range(N)]},
 'target_history': {'rows':trend}, 'fishbone': {'matched':True,'defect':'D0','causes':causes,'rca':rca}, 'fishbone_style':{}
}

with tempfile.TemporaryDirectory() as td:
    out=Path(td)
    xlsx=out/'report.xlsx'; pdf=out/'report.pdf'; pptx=out/'report.pptx'
    xlsx.write_bytes(_excel_report(payload)); pdf.write_bytes(_pdf_report(payload)); pptx.write_bytes(_pptx_report(payload))
    wb=load_workbook(xlsx,read_only=True,data_only=False)
    expected={'Defect Analysis':N+4,'Work Center':N+4,'Grade Analysis':N+4,'Monthly Trend':N+3,'Weekly Trend':N+3,'Quarterly Trend':12+1,'Financial Year':7+1,'6M Fishbone Analysis':1+1+1+N*0+48,'Target vs Actual History':N+3}
    for sheet, min_rows in expected.items():
        assert wb[sheet].max_row >= min_rows, f'{sheet} appears truncated: {wb[sheet].max_row}'
    prs=Presentation(pptx)
    required=('Quality Decision Distribution','Defect Analysis','Work Center Performance','Grade Performance','Monthly Trend','Weekly Trend','Quarterly Trend','Financial Year Trend')
    hits={k:0 for k in required}
    for slide in prs.slides:
        text=' | '.join(sh.text for sh in slide.shapes if hasattr(sh,'text'))
        pics=sum(1 for sh in slide.shapes if getattr(sh,'shape_type',None)==13)
        tables=sum(1 for sh in slide.shapes if getattr(sh,'has_table',False))
        for k in required:
            if k in text:
                hits[k]+=1
                assert pics >= 1, f'{k} slide missing chart image'
                assert tables >= 1, f'{k} slide missing paired table'
    for k,v in hits.items():
        assert v >= 1, f'{k} missing from PPT'
print('EXPORT ACCEPTANCE PASS')
