import os, sys, time
from pathlib import Path

ROOT = Path(__file__).resolve().parent
if not (ROOT / 'server.py').exists():   # works from the repo root or from a tests/ subfolder
    ROOT = ROOT.parent
sys.path.insert(0, str(ROOT))
os.environ['APP_VERSION'] = 'test'

from reports import _excel_report, _pdf_report, _pptx_report

def mk_rows(n, prefix):
    return [{'name':f'{prefix}_{i}','output_qty':float(i+1),'coils':1,'defect_coils':0,'defect_pct':0.0,'reject_qty':0.0,'reject_pct_qty':0.0} for i in range(n)]

N=2000
wc=mk_rows(N,'WC'); gr=mk_rows(N,'GR')
trend=[{'name':f'P{i}','coils':100,'output_qty':100.0,'defect_coils':5,'defect_pct':0.05,'reject_qty':3.0,'reject_pct_qty':0.03,'first_pass_yield_pct':0.92} for i in range(120)]
reg=[{'rank':i+1,'defect':f'D{i}','records':1,'qty':1.0,'pct_records':0.001,'cum_pct':min(1,(i+1)/10)} for i in range(N)]
causes={k:[f'Cause {i}' for i in range(200)] for k in ('man','machine','material','method','measurement','environment')}
rca={k:[{'root_cause':f'RC{i}','action':'Action','preventive_action':'Prevent','role':'Role','responsibility':'Owner'} for i in range(100)] for k in causes}

payload={
 'filters': {'month':'All','week':'All','quarter':'All','financial_year':'All','work_center':'All','grade':'All','main_defect':'All','defect_intensity':'All','quality_decision':'All'},
 'kpis': {'kpis':[{'label':f'K{i}','value':i,'fmt':''} for i in range(12)],'decision_table':[{'decision':'PRIME','qty':100.0}],'intensity_table':[{'intensity':'LIGHT','qty':100.0,'coils':10}]},
 'defects': {'register':reg,'pareto':reg[:10],'register_total':{'records':N,'qty':float(N),'pct_records':1.0}},
 'wcg': {'by_work_center':wc,'by_grade':gr,'total_work_center':{},'total_grade':{}},
 'monthly': {'rows':trend[:100],'total':{}}, 'period': {'rows':trend[:100],'total':{}}, 'quarterly': {'rows':trend[:40],'total':{}}, 'yearly': {'rows':trend[:20],'total':{}},
 'intel': {'early_warnings':[{'title':'Warning','detail':'High cardinality stress','action':'Review'}]*60,'recurring_patterns':[{'defect':'D','grade':'G','work_center':'W','period_count':5,'qty':1.0}]*100,'kpi_ranking':[],'risk_matrix':{'work_centers':wc,'grades':gr},'health_score':{'reasons':[]}},
 'root_cause': {'defect':'D0','rows':[{'grade':'A','work_center':'WC','heat_no':'H','batch_no':'B','output_weight':1.0} for _ in range(1000)]},
 'target_history': {'rows':trend[:100]}, 'fishbone': {'matched':True,'defect':'D0','causes':causes,'rca':rca}, 'fishbone_style':{}
}

import io, shutil, subprocess, tempfile
outputs={}
for name, fn in [('excel', _excel_report), ('pdf', _pdf_report), ('pptx', _pptx_report)]:
    t=time.time(); data=fn(payload); assert data and len(data)>1000
    outputs[name]=data
    print(f'{name}_export=PASS bytes={len(data)} sec={time.time()-t:.2f}')

# Completeness: high-cardinality tables must be printed IN FULL (no top-N cap), only chart visuals are bounded.
import openpyxl
wb=openpyxl.load_workbook(io.BytesIO(outputs['excel']), read_only=True)
for sheet in ('Defect Analysis','Work Center','Grade Analysis'):
    n=sum(1 for _ in wb[sheet].iter_rows(values_only=True))
    assert n>=N, f'Excel sheet {sheet} has only {n} rows (< {N})'
from pptx import Presentation
prs=Presentation(io.BytesIO(outputs['pptx']))
totals={}
for slide in prs.slides:
    title=next((sh.text_frame.text.strip().split('\n')[0] for sh in slide.shapes if getattr(sh,'has_text_frame',False) and sh.text_frame.text.strip()), None)
    for sh in slide.shapes:
        if getattr(sh,'has_table',False) and sh.has_table:
            totals[title]=totals.get(title,0)+len(sh.table.rows)-1
for title in ('Defect Analysis','Work Center Performance','Grade Performance'):
    assert totals.get(title,0)>=N, f'PPTX "{title}" lists only {totals.get(title,0)} rows (< {N})'
if shutil.which('pdftotext'):
    fp=tempfile.mktemp(suffix='.pdf'); open(fp,'wb').write(outputs['pdf'])
    text=subprocess.run(['pdftotext','-layout',fp,'-'],capture_output=True,text=True).stdout
    for tag in (f'D{N-1}', f'WC_{N-1}', f'GR_{N-1}'):
        assert tag in text, f'PDF is missing the last row {tag} (table truncated)'
    assert 'Showing top' not in text, 'PDF still contains a top-N truncation note'
print('EXPORT STRESS PASS')
